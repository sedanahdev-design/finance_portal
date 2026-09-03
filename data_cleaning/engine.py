"""محرك تنظيف الداتا (الفكرة الثالثة).

يحوّل أربعة مصادر خام إلى صيغة موحّدة نظيفة تطابق شكل الملف المرجعي
(تجميع معدل - النتيجة): الفاتورة / اسم الزبون / مركز الكلفة / القيمة
المؤجلة / تاريخ التسليم / الكتلة.

المصادر:
    1) حركة إجمالي الفواتير - مبيعات  → الفاتورة = "مبيع"
    2) حركة إجمالي الفواتير - مرتجعات → الفاتورة = "مرتجع"، والقيمة سالبة
    3) دفتر أستاذ الحسم الممنوح       → جدول مستقل (رقم السند="دفعة"،
                                          مدين سالب، مركز الكلفة، المقابل)
    4) دفاتر أستاذ الذمم (لدى قسم التوزيع / محصلة غير مقبوضة)
                                       → الفاتورة = "تسديد" أو "رصيد سابق"

ملاحظة مهمة (موضّحة للمستخدم في الواجهة): عمود "الكتلة" في الملف المرجعي
هو منطقة/خط جغرافي يُشتق من جدول ربط (اسم المندوب ← المنطقة) غير موجود
ضمن الملفات المرفوعة لهذه الفكرة تحديداً. إن توفّر هذا الجدول لاحقاً يمكن
رفعه اختيارياً لتعبئة هذا العمود تلقائياً؛ وإلى حين توفره يبقى فارغاً بدل
تخمين قيمته.
"""

import re
from decimal import Decimal

import openpyxl

CODE_PREFIX = re.compile(r"^\s*\d+\s*-\s*")


def _clean(v):
    if v is None:
        return ""
    return str(v).strip() if isinstance(v, str) else v


def _to_decimal(v):
    if v is None or v == "":
        return Decimal("0")
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v)).quantize(Decimal("0.01"))
    s = str(v).replace(",", "").strip()
    m = re.match(r"^-?\d+(\.\d+)?", s)
    if not m:
        return Decimal("0")
    try:
        return Decimal(m.group(0)).quantize(Decimal("0.01"))
    except Exception:  # noqa: BLE001
        return Decimal("0")


def strip_code(text):
    """يحذف أي رمز/رقم بادئ من نص مركز الكلفة أو الحساب المقابل، ويبقي النص فقط."""
    text = _clean(text)
    if not text:
        return ""
    text = str(text)
    return CODE_PREFIX.sub("", text).strip()


_ONLY_DIGITS_OR_SYMBOLS = re.compile(r"^[\s0-9.,:/\\\-]*$")
_EMBEDDED_DIGITS_RUN = re.compile(r"[0-9]+")


def safe_cost_center_fallback(value):
    """يُستخدم فقط عندما يكون عمود "مركز الكلفة" فارغاً في الصف الخام (حالة
    موجودة فعلياً في بعض ملفات دفاتر أستاذ الذمم) ونحتاج نصاً بديلاً معقولاً
    من عمود "البيان". لا يُقبل هنا أي رقم صرف (مثل رقم سند/إيصال) ولا أي نص
    يحتوي أرقاماً — إما نص اسم نظيف بالكامل، أو حقل فارغ. هذا أدق من إدراج
    أرقام/رموز في عمود مركز الكلفة، والذي كان يُنتج نتائج غير صحيحة."""
    text = _clean(value)
    if not text:
        return ""
    text = str(text)
    if _ONLY_DIGITS_OR_SYMBOLS.match(text):
        return ""  # قيمة رقمية صرفة (رقم سند مثلاً) — لا تصلح كمركز كلفة
    if _EMBEDDED_DIGITS_RUN.search(text):
        return ""  # نص وصفي طويل يحوي أرقاماً (تاريخ/رقم إيصال) — ليس اسم مركز كلفة فعلي
    return text.strip()


def find_header(all_rows, required_cols, max_scan=10):
    for i, row in enumerate(all_rows[:max_scan]):
        cells = {str(_clean(c)) for c in row if c is not None}
        if required_cols.issubset(cells):
            return i, [str(_clean(c)) for c in row]
    return None, None


def _col_map(header):
    return {h: header.index(h) for h in header if h}


def parse_sales_or_returns(file_obj, is_return):
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header_idx, header = find_header(all_rows, {"الفاتورة", "اسم الزبون", "مركز الكلفة", "القيمة المؤجلة"})
    if header_idx is None:
        raise ValueError("لم يتم العثور على صف العناوين المتوقع (الفاتورة / اسم الزبون / مركز الكلفة / القيمة المؤجلة).")
    col = _col_map(header)

    def get(row, name):
        idx = col.get(name)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    out = []
    for row in all_rows[header_idx + 1:]:
        if row is None or all(c in (None, "") for c in row):
            continue
        first_cell = _clean(row[0]) if row[0] is not None else ""
        if first_cell == "الفاتورة":
            # صف عناوين ثانٍ يبدأ قسم "المجموع/مجموع الآجل/مجموع النقدي" —
            # أي أن كل ما بعده صفوف مجاميع فرعية وإجمالي عام حسب التصنيف
            # (مثل "عرض حليب و معقمات"، "الإجمالي"...) وليست حركات فعلية،
            # وكانت تُقرأ خطأً كصفوف بيانات (مع تسريب أرقام إلى عمود مركز
            # الكلفة لأن أعمدتها بترتيب مختلف تماماً عن صفوف الحركات).
            break
        customer = _clean(get(row, "اسم الزبون"))
        if not customer:
            continue
        deferred = _to_decimal(get(row, "القيمة المؤجلة"))
        if is_return and deferred > 0:
            deferred = -deferred
        out.append({
            "invoice": "مرتجع" if is_return else "مبيع",
            "customer": customer,
            "cost_center": strip_code(get(row, "مركز الكلفة")),
            "deferred_value": deferred,
            "delivery_date": _clean(get(row, "تاريخ التسليم")),
            "zone": "",
            "source": "مرتجعات" if is_return else "مبيعات",
        })
    return out


def parse_discount_ledger(file_obj):
    """يقرأ دفتر أستاذ الحسم الممنوح ويعيد جدول الدفعات (رقم السند / مدين سالب / تاريخ / مركز الكلفة / المقابل)."""
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header_idx, header = find_header(all_rows, {"رقم السند", "مدين", "دائن", "التاريخ"})
    if header_idx is None:
        raise ValueError("لم يتم العثور على صف العناوين المتوقع داخل دفتر الحسم الممنوح.")
    col = _col_map(header)

    def get(row, name):
        idx = col.get(name)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    out = []
    for row in all_rows[header_idx + 1:]:
        if row is None or all(c in (None, "") for c in row):
            continue
        first_cell = _clean(row[0])
        if str(first_cell) in ("المجموع", "الرصيد النهائي"):
            break  # نهاية القيود الفعلية، وما بعدها صفوف مجاميع/رصيد فقط
        if first_cell == "الرصيد السابق":
            continue
        debit = _to_decimal(get(row, "مدين"))
        raw_date = get(row, "التاريخ")
        if debit == 0 or not raw_date:
            continue
        cost_center = strip_code(get(row, "مركز الكلفة"))
        counterpart = strip_code(get(row, "الحساب المقابل"))
        out.append({
            "voucher_no": "دفعة",
            "debit": -abs(debit),
            "date": _clean(raw_date),
            "cost_center": cost_center or counterpart,
            "counterpart": counterpart,
            "narration": _clean(get(row, "البيان")),
        })
    return out


def parse_receivables_ledger(file_obj):
    """يقرأ دفتر أستاذ ذمم (لدى قسم التوزيع / محصلة غير مقبوضة) ويحوّله لصيغة موحّدة (تسديد / رصيد سابق)."""
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header_idx, header = find_header(all_rows, {"رقم السند", "مدين", "دائن", "التاريخ"})
    if header_idx is None:
        raise ValueError("لم يتم العثور على صف العناوين المتوقع داخل دفتر ذمم.")
    col = _col_map(header)

    def get(row, name):
        idx = col.get(name)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    out = []
    current_subaccount = ""
    for row in all_rows[header_idx + 1:]:
        if row is None or all(c in (None, "") for c in row):
            continue
        first_cell = _clean(row[0])
        narration = _clean(get(row, "البيان"))
        raw_date = get(row, "التاريخ")
        credit = _to_decimal(get(row, "دائن"))
        debit = _to_decimal(get(row, "مدين"))

        if str(first_cell) in ("المجموع", "الرصيد النهائي"):
            break

        is_header_or_balance = str(first_cell) in ("0", "") and (raw_date in (0, "0", None, "") or "-" not in str(raw_date))
        if is_header_or_balance:
            if narration == "الرصيد السابق":
                net = credit - debit
                if net != 0:
                    out.append({
                        "invoice": "رصيد سابق", "customer": strip_code(current_subaccount),
                        "cost_center": "", "deferred_value": net, "delivery_date": "", "zone": "",
                        "source": "ذمم",
                    })
            elif narration and isinstance(narration, str) and re.match(r"^\d+-", narration):
                current_subaccount = narration
            continue

        if credit == 0 and debit == 0:
            continue
        counterpart = strip_code(get(row, "الحساب المقابل"))
        cost_center = strip_code(get(row, "مركز الكلفة"))
        out.append({
            "invoice": "تسديد",
            "customer": counterpart or strip_code(current_subaccount),
            "cost_center": cost_center or safe_cost_center_fallback(narration),
            "deferred_value": credit if credit else -debit,
            "delivery_date": _clean(raw_date) if str(raw_date) not in ("0", "") else "",
            "zone": "",
            "source": "ذمم",
        })
    return out


def combine(sales_rows, returns_rows, receivables_rows_list):
    combined = list(sales_rows) + list(returns_rows)
    for rows in receivables_rows_list:
        combined.extend(rows)
    return combined


def summarize(combined_rows, discount_rows):
    total_deferred = sum((r["deferred_value"] for r in combined_rows), Decimal("0"))
    by_source = {}
    for r in combined_rows:
        by_source.setdefault(r["source"], 0)
        by_source[r["source"]] += 1
    return {
        "total_rows": len(combined_rows),
        "by_source": by_source,
        "total_deferred_value": total_deferred,
        "discount_rows": len(discount_rows),
        "discount_total": sum((r["debit"] for r in discount_rows), Decimal("0")),
    }
