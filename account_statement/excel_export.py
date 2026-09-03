import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    headers = ["التاريخ", "رقم السند", "البيان", "مدين", "دائن", "العملة كما وردت", "الفئة"]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for r in rows:
        ws.append([r["date"], r["voucher_no"], r["narration"], float(r["debit"]), float(r["credit"]),
                   r["currency_raw"], r["category"]])
    widths = [12, 12, 50, 16, 16, 26, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def build_workbook(result, meta):
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "الملخص"
    ws0.sheet_view.rightToLeft = True
    ws0.append(["كشف حساب:", meta.get("file_name", "")])
    ws0.append(["إجمالي عدد الحركات:", result["total_rows"]])
    ws0.append([])
    ws0.append([
        "العملة", "عدد الحركات", "إجمالي مدين (بعملته الأصلية)", "إجمالي دائن (بعملته الأصلية)",
        "الرصيد (بعملته الأصلية)", "إجمالي مدين (معادل ل.س.ج)", "إجمالي دائن (معادل ل.س.ج)",
    ])
    for cell in ws0[4]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for code, s in result["summary"].items():
        ws0.append([
            s["label"], s["count"], float(s["total_debit"]), float(s["total_credit"]), float(s["balance"]),
            float(s.get("total_debit_syp_equivalent", s["total_debit"])),
            float(s.get("total_credit_syp_equivalent", s["total_credit"])),
        ])
    ws0.append([])
    ws0.append([
        "ملاحظة: عمودا \"مدين/دائن\" في كل شيت عملة يُظهران القيمة الحقيقية بعملة"
        " الحركة الأصلية (دولار/ليرة قديمة/ليرة جديدة) لا القيمة المحوَّلة لليرة"
        " الجديدة كما وردت في الملف الخام — العمودان الأخيران هنا للمراجعة"
        " والتأكد من مطابقة المجموع لإجمالي كشف الحساب الأصلي غير المُقسَّم."
    ])
    for i, w in enumerate([22, 14, 22, 22, 20, 20, 20], start=1):
        ws0.column_dimensions[get_column_letter(i)].width = w

    for code, bucket_rows in result["buckets"].items():
        from account_statement.engine import CURRENCY_LABELS
        _write_sheet(wb, CURRENCY_LABELS.get(code, code), bucket_rows)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
