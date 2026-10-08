import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _header(ws, headers):
    ws.append(headers)
    r = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = f"A{r + 1}"


def _write_ficha_group(ws, col_start, rows, value_key, extra1_key, extra2_key):
    """يكتب مجموعة أعمدة واحدة من فيشة 'كشف الحساب' (قيمة الكشوفات /
    الدفعات / مرتجعات) بدءاً من الصف 4 في العمود col_start، ثم صف
    'الإجمالي' مباشرة بعد آخر صف بيانات لهذه المجموعة بالذات (بصرف النظر
    عن عدد صفوف المجموعات الأخرى) — طبقاً للفيشة المرجعية الحقيقية (شادي
    حسن، 2026-10-08) حيث يختلف موضع صف 'الإجمالي' باختلاف طول كل مجموعة."""
    total = 0.0
    for i, r in enumerate(rows):
        row_no = 4 + i
        ws.cell(row=row_no, column=col_start, value=float(r[value_key]))
        ws.cell(row=row_no, column=col_start + 1, value=r[extra1_key])
        ws.cell(row=row_no, column=col_start + 2, value=r[extra2_key])
        total += float(r[value_key])
    total_row = 4 + len(rows)
    ws.cell(row=total_row, column=col_start, value=total)
    ws.cell(row=total_row, column=col_start + 1, value="الإجمالي")
    if len(rows) > 0:
        ws.merge_cells(start_row=total_row, start_column=col_start + 1,
                        end_row=total_row, end_column=col_start + 2)
    return total_row


def _build_ficha_sheet(wb, d):
    """شيت 'فيشة' إضافي لكل موزّع — يكرّر حرفياً صيغة الفيشة اليدوية
    الحقيقية التي رفعها المستخدم (كشف حساب السيد شادي حسن شهر 9،
    2026-10-08)، بالإضافة لشيت المدين/دائن الموجود مسبقاً (لا يستبدله،
    بطلب المستخدم الصريح). راجع توثيق الاجتهادات غير المؤكدة بعد (رقم
    الايصال/اسم الصيدلية/تسميات التواريخ) في رأس external_commissions/engine.py."""
    sheet_name = f"{d['name'][:22]}-فيشة" if d["name"] else "فيشة"
    ws = wb.create_sheet(sheet_name[:31])
    ws.sheet_view.rightToLeft = True

    month_label = f" شهر {d['month_number']}" if d["month_number"] else ""
    ws.append([f"كشف حساب السيد {d['name']}{month_label}"])
    ws.merge_cells("A1:K1")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A1"].font = Font(bold=True)

    ws.append(["قيمة الكشوفات", None, None, None, "الدفعات", None, None, None, "مرتجعات"])
    ws.merge_cells("A2:C2")
    ws.merge_cells("E2:G2")
    ws.merge_cells("I2:K2")
    for coord in ("A2", "E2", "I2"):
        ws[coord].font = HEADER_FONT
        ws[coord].fill = HEADER_FILL
        ws[coord].alignment = Alignment(horizontal="center")

    ws.append(["قيمة الكشف", "رقم الكشف", "التاريخ", None, "الدفعة", "التاريخ", "رقم الايصال",
               None, "الدفعة", "التاريخ", "اسم الصيدلية"])
    for col_idx in (1, 2, 3, 5, 6, 7, 9, 10, 11):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    debit_rows = [{"amount": r["amount"], "voucher_no": r["voucher_no"], "date": r["date"]}
                  for r in d["debit_rows"]]
    payment_rows = [r for r in d["credit_rows"] if not r["is_return"]]
    return_rows = [r for r in d["credit_rows"] if r["is_return"]]

    totals_rows = [
        _write_ficha_group(ws, 1, debit_rows, "amount", "voucher_no", "date"),
        _write_ficha_group(ws, 5, payment_rows, "amount", "date", "receipt_no"),
        _write_ficha_group(ws, 9, return_rows, "amount", "date", "pharmacy_name"),
    ]
    summary_start = max(totals_rows) + 3

    opening_label = f"قيمة الرصيد بتاريخ{d['opening_date_label']}" if d["opening_date_label"] else "قيمة الرصيد الافتتاحي"
    closing_label = f"رصيد{d['closing_date_label']}" if d["closing_date_label"] else "الرصيد الختامي"
    returns_label = f"مرتجعات شهر{d['month_number']}" if d["month_number"] else "مرتجعات الشهر"

    ws.cell(row=summary_start, column=1, value=opening_label)
    ws.cell(row=summary_start, column=2, value=float(d["opening_balance"]))
    ws.cell(row=summary_start + 1, column=1, value="قيمة الكشوفات")
    ws.cell(row=summary_start + 1, column=2, value=float(d["total_debit"]))
    ws.cell(row=summary_start + 2, column=1, value=f"الدفعات المستلمة من السيد {d['name']}")
    ws.cell(row=summary_start + 2, column=2, value=float(d["total_payments"]))
    ws.cell(row=summary_start + 3, column=1, value=returns_label)
    ws.cell(row=summary_start + 3, column=2, value=float(d["total_returns"]))
    ws.cell(row=summary_start + 4, column=1, value=closing_label)
    ws.cell(row=summary_start + 4, column=2, value=float(d["closing_balance"]))

    commission_row = summary_start + 6
    ws.cell(row=commission_row, column=1, value="العمولة =(الكشوفات - مرتجعات )*3%")
    ws.cell(row=commission_row, column=6,
            value=f"({float(d['total_debit']):,.0f}-{float(d['total_returns']):,.0f})*3%")
    ws.cell(row=commission_row + 1, column=1, value=float(d["commission"]))

    for i, w in enumerate([14, 12, 12, 3, 14, 12, 14, 3, 14, 12, 26], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def build_workbook(distributors, summary):
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "الملخص"
    ws0.sheet_view.rightToLeft = True
    ws0.append(["معادلة العمولة الحالية: (كشوفات − مرتجعات) × 3%  — تعديل 2026-10-08 بطلب "
                "المستخدم الصريح طبقاً لفيشة مرجعية حقيقية (شادي حسن)؛ عوّضت المعادلة السابقة "
                "(مرتجع − كشوفات) × 3% بتاريخ 2026-08-26، وتلك عوّضت الأصلية "
                "(مبيع + مرتجع − كشوفات) × 3%."])
    ws0.append([])
    _header(ws0, ["الموزّع", "الرصيد الافتتاحي (دوّار)", "إجمالي الكشوفات (مدين)",
                  "إجمالي الدائن (دفعات + مرتجعات)", "منها: مرتجعات (أساس العمولة)",
                  "منها: دفعات (مستبعدة من العمولة)", "أساس العمولة (كشوفات − مرتجعات)",
                  "نسبة العمولة", "العمولة", "الرصيد الختامي"])
    for d in distributors:
        ws0.append([
            d["name"], float(d["opening_balance"]), float(d["total_debit"]), float(d["total_credit"]),
            float(d["total_returns"]), float(d["total_payments"]),
            float(d["commission_base"]), f"{float(d['commission_rate'])*100:.0f}%",
            float(d["commission"]), float(d["closing_balance"]),
        ])
    ws0.append([])
    ws0.append(["الإجمالي", "", float(summary["total_debit"]), float(summary["total_credit"]),
                float(summary["total_returns"]), float(summary["total_payments"]), "", "",
                float(summary["total_commission"]), ""])
    for i, w in enumerate([26, 20, 20, 20, 18, 22, 20, 14, 16, 18], start=1):
        ws0.column_dimensions[get_column_letter(i)].width = w

    for d in distributors:
        ws = wb.create_sheet(d["name"][:28] or "موزّع")
        ws.sheet_view.rightToLeft = True
        ws.append([f"كشف حساب: {d['name']}"])
        ws.append(["الرصيد الافتتاحي (دوّار من الشهر السابق)", float(d["opening_balance"])])
        ws.append([])
        ws.append(["العمولة = (كشوفات − مرتجعات) × 3%  — عمود \"النوع\" يوضّح أي حركة دائن اعتُبرت "
                    "مرتجعاً (تدخل في أساس العمولة) وأيها دفعة (مستبعدة من العمولة، وليست محذوفة)."])
        ws.append(["قيمة الكشوفات (مدين)", None, None, "الدفعات + المرتجعات (دائن)"])
        _header(ws, ["رقم السند", "المرجع", "التاريخ", "المبلغ (مدين)", "",
                     "رقم السند", "المرجع", "التاريخ", "المبلغ (دائن)", "النوع"])
        max_len = max(len(d["debit_rows"]), len(d["credit_rows"]))
        for i in range(max_len):
            deb = d["debit_rows"][i] if i < len(d["debit_rows"]) else None
            cre = d["credit_rows"][i] if i < len(d["credit_rows"]) else None
            ws.append([
                deb["voucher_no"] if deb else "", deb["reference"] if deb else "",
                deb["date"] if deb else "", float(deb["amount"]) if deb else "",
                "",
                cre["voucher_no"] if cre else "", cre["reference"] if cre else "",
                cre["date"] if cre else "", float(cre["amount"]) if cre else "",
                cre["kind"] if cre else "",
            ])
        ws.append([])
        ws.append(["إجمالي مدين (كشوفات)", float(d["total_debit"]), "", "إجمالي دائن (دفعات + مرتجعات)", float(d["total_credit"])])
        ws.append(["منها: مرتجعات (أساس العمولة)", float(d["total_returns"]), "", "منها: دفعات (مستبعدة من العمولة)", float(d["total_payments"])])
        ws.append(["أساس العمولة (كشوفات − مرتجعات)", float(d["commission_base"])])
        ws.append(["العمولة (3%)", float(d["commission"])])
        ws.append(["الرصيد الختامي", float(d["closing_balance"])])
        for i, w in enumerate([12, 30, 12, 14, 4, 12, 30, 12, 14, 22], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

        _build_ficha_sheet(wb, d)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
