import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TIER1_FILL = PatternFill("solid", fgColor="E7F7EF")       # أخضر فاتح
TIER2_FILL = PatternFill("solid", fgColor="E0F2FE")       # أزرق فاتح
TIER3_FILL = PatternFill("solid", fgColor="FEF3C7")       # كهرماني فاتح
TIER4_FILL = PatternFill("solid", fgColor="EDE9FE")       # بنفسجي فاتح
BAD_FILL = PatternFill("solid", fgColor="FDEAEA")         # أحمر فاتح


def _f(v):
    if isinstance(v, Decimal):
        return float(v)
    return v


def _header(ws, headers):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _write_price_match_sheet(wb, title, match_rows, fill, vita_compliance=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    headers = [
        "الفاتورة", "التاريخ", "اسم الزبون", "اسم المادة", "الكمية", "الهدايا",
        "الإفرادي الفعلي", "أساس المقارنة", "النطاق المتوقع (من)", "النطاق المتوقع (إلى)",
        "عرض المفرق", "عرض مميز1", "عرض مميز",
    ]
    show_vita = vita_compliance is not None
    if show_vita:
        headers += ["فحص فيتا فارما: مطالبة الهدايا الصافية للحزمة (مادة×عرض مفرق)", "التزام العروض"]
    _header(ws, headers)
    for m in match_rows:
        r = m.row
        row = [
            r.invoice, r.date, r.customer, r.item, _f(r.qty), _f(r.gifts),
            _f(r.unit_price), m.basis, _f(m.expected_lo), _f(m.expected_hi),
            r.retail_offer, r.special_offer1, r.special_offer,
        ]
        if show_vita:
            g = vita_compliance.get((r.item, r.retail_offer))
            claim = g.claim_value if g is not None else None
            row += [_f(claim) if claim is not None else "لا ينطبق (بلا عرض مفرق صالح)",
                    ("ملتزم" if claim == 0 else "غير ملتزم (انحراف بالهدايا)") if claim is not None else "—"]
        ws.append(row)
        for c in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=c).fill = fill
    widths = [16, 12, 30, 40, 9, 9, 13, 30, 15, 15, 12, 12, 12, 30, 20]
    for i, w in enumerate(widths[:len(headers)], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    return ws


DOLLAR_LABEL = {True: "نعم", False: "لا", None: "غير مؤكَّد"}


def _write_unmatched_sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    headers = [
        "الفاتورة", "التاريخ", "اسم الزبون", "اسم المادة", "الكمية", "الهدايا",
        "الإفرادي", "السعر الإجمالي", "الحسم", "حسم هدايا",
        "عرض المفرق", "عرض مميز1", "عرض مميز", "دولاري؟", "ملاحظة العملة", "ملاحظة",
    ]
    _header(ws, headers)
    for r in rows:
        note = "كمية = صفر — السعر الإفرادي هنا غير ذي دلالة (لا معاملة فعلية بهذا السطر)" if r.qty == 0 else ""
        ws.append([
            r.invoice, r.date, r.customer, r.item, _f(r.qty), _f(r.gifts),
            _f(r.unit_price), _f(r.total_price), _f(r.discount), _f(r.gift_discount),
            r.retail_offer, r.special_offer1, r.special_offer,
            DOLLAR_LABEL[r.is_dollar], r.dollar_note, note,
        ])
        for c in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=c).fill = BAD_FILL
    widths = [16, 12, 30, 40, 9, 9, 12, 14, 10, 10, 12, 12, 12, 10, 45, 55]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    return ws


def _disc_ratio_f(row):
    if row.total_price:
        return float((row.discount / row.total_price).quantize(Decimal("0.0001")))
    return 0.0


def _write_reconciliation_sheets(wb, title_matched, title_mismatched, title_no_counterpart, reconciliation, with_currency=True):
    def _write(ws_title, groups, fill=None):
        ws = wb.create_sheet(ws_title)
        ws.sheet_view.rightToLeft = True
        headers = ["اسم الزبون", "اسم المادة", "عدد أسطر بيع", "عدد أسطر مرتجع",
                   "أسعار البيع", "أسعار المرتجع", "نسب حسم البيع", "نسب حسم المرتجع"]
        if with_currency:
            headers += ["دولاري؟ (بيع)", "دولاري؟ (مرتجع)"]
        headers += ["الأسباب / الملاحظات"]
        _header(ws, headers)
        for g in groups:
            sale_prices = sorted({float(s.unit_price) for s in g.sale_rows})
            return_prices = sorted({float(s.unit_price) for s in g.return_rows})
            row = [g.customer, g.item, len(g.sale_rows), len(g.return_rows),
                   ", ".join(map(str, sale_prices)), ", ".join(map(str, return_prices)),
                   ", ".join(map(str, sorted({_disc_ratio_f(s) for s in g.sale_rows}))),
                   ", ".join(map(str, sorted({_disc_ratio_f(s) for s in g.return_rows})))]
            if with_currency:
                row += [
                    ", ".join(sorted({DOLLAR_LABEL[getattr(s, "is_dollar", None)] for s in g.sale_rows})) if g.sale_rows else "—",
                    ", ".join(sorted({DOLLAR_LABEL[getattr(s, "is_dollar", None)] for s in g.return_rows})) if g.return_rows else "—",
                ]
            row += [" | ".join(g.reasons) if g.reasons else "متطابق بالسعر والحسم" + (" والعملة" if with_currency else "")]
            ws.append(row)
            if fill is not None:
                for c in range(1, len(headers) + 1):
                    ws.cell(row=ws.max_row, column=c).fill = fill
        widths = [30, 40, 12, 12, 20, 20, 16, 16, 12, 14, 50]
        for i, w in enumerate(widths[:len(headers)], start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        return ws

    _write(title_matched, reconciliation["matched"])
    _write(title_mismatched, reconciliation["mismatched"], fill=BAD_FILL)
    _write(title_no_counterpart, reconciliation["no_counterpart"])


def build_workbook(result, meta):
    wb = Workbook()
    s = result["summary"]
    rate_range = result["rate_range"]
    price_data = result["price_data"]
    ledger_data = result["ledger_data"]

    ws0 = wb.active
    ws0.title = "الملخص"
    ws0.sheet_view.rightToLeft = True
    rows = [
        ["ملف الحركة اليومية (مفرق)", meta.get("movement_file_name", "")],
        ["ملف قائمة الأسعار بالدولار", meta.get("price_file_name", "")],
        ["ملف دفتر الأستاذ (مواد دولار — لمعرفة العملة فقط)", meta.get("ledger_file_name", "")],
        ["ملف مبيعات الجملة مع المرتجعات", meta.get("wholesale_file_name", "") or "لم يُرفَع"],
        ["", ""],
        ["نمط سعر الصرف", "ثابت" if rate_range.lo == rate_range.hi else "متغير (نطاق)"],
        ["سعر الصرف المستخدَم (من)", float(rate_range.lo)],
        ["سعر الصرف المستخدَم (إلى)", float(rate_range.hi)],
        ["", ""],
        ["عدد مواد قائمة الأسعار الصالحة (سعر > 0)", len(price_data["prices"])],
        ["مواد سعرها صفر بالقائمة (دعائية، استُبعدت)", ", ".join(price_data["zero_priced"]) or "لا يوجد"],
        ["عدد قيود دفتر الأستاذ الدولارية المُستخرَجة", ledger_data["parsed_count"]],
        ["", ""],
        ["إجمالي أسطر الحركة اليومية", s["movement_rows_count"]],
        ["مطابقين مباشرة (المرحلة 1)", s["tier1_count"]],
        ["مطابقين بعد خصم الحسم (المرحلة 2)", s["tier2_count"]],
        ["مطابق على العرض المفرق (المرحلة 3)", s["tier3_count"]],
        ["مطابق على العرض المميز (المرحلة 4)", s["tier4_count"]],
        ["غير مطابقين بكل الطرق السابقة", s["unmatched_count"]],
        ["", ""],
        ["مطابقة مبيعات/مرتجعات مفرق — متطابق", s["retail_matched_count"]],
        ["مطابقة مبيعات/مرتجعات مفرق — غير متطابق (بيع ومرتجع معاً لكن يختلفان)", s["retail_mismatched_count"]],
        ["مطابقة مبيعات/مرتجعات مفرق — بلا طرف مقابل (معلوماتي، ليس خطأ)", s["retail_no_counterpart_count"]],
    ]
    if result.get("wholesale_reconciliation") is not None:
        rows += [
            ["مطابقة مبيعات/مرتجعات جملة — متطابق", s["wholesale_matched_count"]],
            ["مطابقة مبيعات/مرتجعات جملة — غير متطابق (بيع ومرتجع معاً لكن يختلفان)", s["wholesale_mismatched_count"]],
            ["مطابقة مبيعات/مرتجعات جملة — بلا طرف مقابل (معلوماتي، ليس خطأ)", s["wholesale_no_counterpart_count"]],
        ]
    for row in rows:
        ws0.append(row)
    ws0.column_dimensions["A"].width = 48
    ws0.column_dimensions["B"].width = 40

    pm = result["price_match"]
    _write_price_match_sheet(wb, "مطابقين", pm.tier1, TIER1_FILL)
    _write_price_match_sheet(wb, "مطابقين بعد خصم الحسم", pm.tier2, TIER2_FILL)
    _write_price_match_sheet(wb, "مطابق على العرض المفرق", pm.tier3_retail, TIER3_FILL, vita_compliance=pm.vita_compliance)
    _write_price_match_sheet(wb, "مطابق على العرض المميز", pm.tier4_special, TIER4_FILL)
    _write_unmatched_sheet(wb, "غير مطابقين بكل الطرق السابقة", pm.unmatched)

    _write_reconciliation_sheets(
        wb, "مطابقة مفرق (مبيع⇄مرتجع)", "تنبيه - غير مطابق مفرق", "بلا طرف مقابل - مفرق",
        result["retail_reconciliation"], with_currency=True,
    )
    if result.get("wholesale_reconciliation") is not None:
        _write_reconciliation_sheets(
            wb, "مطابقة جملة (مبيع⇄مرتجع)", "تنبيه - غير مطابق جملة", "بلا طرف مقابل - جملة",
            result["wholesale_reconciliation"], with_currency=False,
        )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
