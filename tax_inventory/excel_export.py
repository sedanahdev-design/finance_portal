import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
NOTE_FILL = PatternFill("solid", fgColor="FFF7E0")


def build_workbook(result, meta, detail_rows=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "نتيجة مطابقة الجرد"
    ws.sheet_view.rightToLeft = True

    headers = [
        "رمز المادة", "اسم المادة", "سقف أمين 9", "إجمالي أمين 8 (حقيقي كامل)",
        "الكمية المسموح بيعها (بعد التوزيع)", "هل تجاوز السقف؟", "عدد الصيدليات",
        "عدد الصيدليات المعتمدة", "ملاحظات",
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = "A2"

    for r in result["rows"]:
        ws.append([
            r["code"], r["name"], r["inventory_qty"], r["sold_qty"], r["final_qty"],
            "نعم" if r["is_capped"] else "لا", r["pharmacy_count"], r["selected_pharmacy_count"], r["note"],
        ])
        if r["is_capped"] or r["note"]:
            for c in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=c).fill = NOTE_FILL

    widths = [14, 46, 14, 20, 24, 14, 14, 16, 46]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # شيت التوزيع التفصيلي على الصيدليات (أساس التوزيع لكل مادة تجاوزت السقف)
    pharmacy_rows = result.get("pharmacy_rows") or []
    if pharmacy_rows:
        ws_ph = wb.create_sheet("التوزيع على الصيدليات")
        ws_ph.sheet_view.rightToLeft = True
        p_headers = ["رمز المادة", "اسم المادة", "الصيدلية (اسم الزبون)", "الكمية (أمين 8)", "الهدايا", "معتمدة ضمن التوزيع؟", "الكمية المعتمدة"]
        ws_ph.append(p_headers)
        for c in range(1, len(p_headers) + 1):
            cell = ws_ph.cell(row=1, column=c)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws_ph.freeze_panes = "A2"
        for r in pharmacy_rows:
            ws_ph.append([
                r["code"], r["material"], r["pharmacy"], r["qty"], r.get("gifts", 0),
                "نعم" if r["selected"] else "لا — مستبعدة من هذا التوزيع", r["allocated_qty"],
            ])
            if not r["selected"]:
                for c in range(1, len(p_headers) + 1):
                    ws_ph.cell(row=ws_ph.max_row, column=c).fill = NOTE_FILL
        p_widths = [14, 46, 40, 16, 12, 26, 18]
        for i, w in enumerate(p_widths, start=1):
            ws_ph.column_dimensions[get_column_letter(i)].width = w

    if detail_rows is not None:
        # نموذج جرد الضريبة التفصيلي: كل حقول حركة المبيعات لمواد (فوق 3 قطع)
        # عدا الفاتورة والتاريخ (بحسب طلب المستخدم صراحةً).
        ws_detail = wb.create_sheet("نموذج جرد الضريبة (تفصيلي)")
        ws_detail.sheet_view.rightToLeft = True
        d_headers = ["اسم الزبون", "اسم المادة", "المجموعة", "كمية", "الهدايا", "اجمالي", "الإفرادي", "السعر الإجمالي"]
        ws_detail.append(d_headers)
        for c in range(1, len(d_headers) + 1):
            cell = ws_detail.cell(row=1, column=c)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws_detail.freeze_panes = "A2"
        for r in detail_rows:
            ws_detail.append([
                r["customer"], r["name"], r["group"], r["qty"], r["gifts"],
                r["total_qty"], r["unit_price"], r["total_price"],
            ])
        d_widths = [46, 46, 22, 10, 10, 10, 14, 16]
        for i, w in enumerate(d_widths, start=1):
            ws_detail.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("الملخص")
    ws2.sheet_view.rightToLeft = True
    s = result["summary"]
    for row in [
        ["ملف جرد أمين 9 (فوق 3 قطع — السقف)", meta.get("over3_file_name", "")],
        ["ملف الجرد (تحت 3 قطع)", meta.get("under3_file_name", "")],
        ["ملف حركة مبيعات أمين 8 (التفصيل الحقيقي حسب الصيدلية)", meta.get("sales_file_name", "")],
        ["", ""],
        ["عدد المواد في جرد أمين 9", s["items_over3"]],
        ["عدد المواد التي وُجد لها مبيعات في أمين 8", s["items_matched_in_sales"]],
        ["إجمالي سقف أمين 9", s["total_inventory_qty"]],
        ["إجمالي أمين 8 الحقيقي الكامل", s["total_sold_qty"]],
        ["إجمالي الكمية المسموح بيعها بعد التوزيع", s["total_final_qty"]],
        ["عدد المواد التي تجاوزت السقف (احتاجت توزيع/أفضل مزيج)", s["capped_items_count"]],
        ["عدد المواد التي استُخدم لها تقريب بدل الحل الأمثل (بيانات ضخمة جداً)", s["approx_items_count"]],
        ["عدد المواد المكرّرة بين القائمتين (فوق/تحت 3)", s["cross_listed_count"]],
    ]:
        ws2.append(row)
    ws2.column_dimensions["A"].width = 42
    ws2.column_dimensions["B"].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
