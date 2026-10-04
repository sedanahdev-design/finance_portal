import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
OK_FILL = PatternFill("solid", fgColor="E7F7EF")
BAD_FILL = PatternFill("solid", fgColor="FDEAEA")
FOUND_DIFF_DATE_FILL = PatternFill("solid", fgColor="FDE9C8")  # برتقالي: موجود بمبلغ مطابق لكن بتاريخ مختلف
NO_DATA_FILL = PatternFill("solid", fgColor="E5E7EB")  # رمادي: لا توجد حركات مقابلة عند هبة (ليس فرقاً رقمياً)
UNRESOLVED_FILL = PatternFill("solid", fgColor="FDE68A")  # كهرماني: مركز كلفة لم يُطابَق بثقة، يحتاج مراجعة يدوية

CC_STATUS_LABEL = {"matched": "مطابق مع هبة", "mismatched": "غير مطابق مع هبة", "no_hiba_data": "لا توجد حركات مقابلة عند هبة"}


def _header(ws, headers):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = "A2"


def build_workbook(result, meta):
    wb = Workbook()
    s = result["summary"]

    ws0 = wb.active
    ws0.title = "الملخص"
    ws0.sheet_view.rightToLeft = True
    for row in [
        ["ملف مشروع دواك", meta.get("dawak_file_name", "")],
        ["ملف كشف حساب هبة", meta.get("hiba_file_name", "")],
        ["", ""],
        ["إجمالي مدين دواك (كل الصيدليات)", float(s["dawak_total_debit"])],
        ["إجمالي دائن دواك (كل الصيدليات)", float(s["dawak_total_credit"])],
        ["", ""],
        ["إجمالي مدين كشف حساب هبة", float(s["hiba_total_debit"])],
        ["إجمالي دائن كشف حساب هبة", float(s["hiba_total_credit"])],
        ["رصيد كشف حساب هبة", float(s["hiba_balance"])],
    ]:
        ws0.append(row)
    ws0.column_dimensions["A"].width = 40
    ws0.column_dimensions["B"].width = 26

    # تصحيح 2026-09-30 (طلب المستخدم الصريح): المطابقة حسب مركز الكلفة هي
    # المقارنة الفعلية الرئيسية مع هبة، فتُبنى وتُلخَّص أولاً — قبل "مطابقة
    # الصيدليات" الداخلية (بون=مرتجع وهمي) التي لا تقارن مع هبة إطلاقاً وقد
    # تُضلِّل لو ظهرت أولاً كأنها "النتيجة".
    cost_centers = result.get("cost_centers")
    if cost_centers is not None:
        ccs = cost_centers["summary"]
        ws0.append(["", ""])
        ws0.append(["مطابقة حسب مركز الكلفة (دواك ⇄ هبة) — المقارنة الفعلية الرئيسية", ""])
        ws0.append(["صيدليات/حسابات مطابقة فعلياً مع هبة", ccs["matched"]])
        ws0.append(["صيدليات/حسابات غير مطابقة فعلياً مع هبة", ccs["mismatched"]])
        ws0.append(["صيدليات لا توجد لها أي حركات هبة مقابلة", ccs["no_hiba_data"]])
        ws0.append(["حركات هبة بمركز كلفة غير مطابَق (تحتاج مراجعة يدوية)", ccs["unresolved_hiba_rows"]])
        ws_cc = wb.create_sheet("مطابقة حسب مركز الكلفة")
        ws_cc.sheet_view.rightToLeft = True
        _header(ws_cc, [
            "رمز الصيدلية", "اسم الصيدلية", "الحالة",
            "مدين دواك", "دائن دواك", "مدين هبة (مُسنَد بمركز الكلفة)", "دائن هبة (مُسنَد بمركز الكلفة)",
            "عدد حركات هبة المُسنَدة", "الفرق: دائن دواك - مدين هبة", "الفرق: مدين دواك - دائن هبة",
        ])
        for r in cost_centers["rows"]:
            ws_cc.append([
                r["code"], r["name"], CC_STATUS_LABEL.get(r["status"], r["status"]),
                float(r["dawak_debit"]), float(r["dawak_credit"]), float(r["hiba_debit"]), float(r["hiba_credit"]),
                r["hiba_rows_count"], float(r["diff_vs_hiba_debit"]), float(r["diff_vs_hiba_credit"]),
            ])
            rr = ws_cc.max_row
            fill = {"matched": OK_FILL, "mismatched": BAD_FILL, "no_hiba_data": NO_DATA_FILL}.get(r["status"], NO_DATA_FILL)
            ws_cc.cell(row=rr, column=3).fill = fill
        for i, w in enumerate([14, 38, 26, 14, 14, 18, 18, 16, 20, 20], start=1):
            ws_cc.column_dimensions[get_column_letter(i)].width = w

        # تصحيح 2026-09-30: تفصيل حركات هبة المُسنَدة لكل صيدلية — حتى يكون
        # أي فرق متبقٍّ (بعد الربط الصحيح بمركز الكلفة) قابلاً للتتبع مباشرة من
        # هنا (أي حركة دواك ناقصة مقابلها عند هبة، أو بالعكس) بلا حاجة للبحث
        # بشيت آخر. طلب المستخدم صراحة توضيح سبب استمرار ظهور "غير مطابق".
        ws_ccd = wb.create_sheet("تفاصيل حركات هبة حسب الصيدلية")
        ws_ccd.sheet_view.rightToLeft = True
        _header(ws_ccd, ["رمز الصيدلية", "اسم الصيدلية", "التاريخ", "رقم القيد", "البيان", "مدين", "دائن"])
        for r in cost_centers["rows"]:
            for e in r["hiba_rows"]:
                ws_ccd.append([r["code"], r["name"], e.get("raw_date", ""), e.get("voucher_no", ""),
                               e.get("narration", ""), float(e.get("debit", 0)), float(e.get("credit", 0))])
        for i, w in enumerate([14, 34, 12, 12, 45, 14, 14], start=1):
            ws_ccd.column_dimensions[get_column_letter(i)].width = w

        # حركات هبة بمركز كلفة لم نجد له أي صيدلية مطابقة بثقة — للمراجعة اليدوية
        # (لا تُسقَط بصمت، حسب فلسفة المشروع الثابتة).
        ws_ccu = wb.create_sheet("مركز كلفة غير مطابق")
        ws_ccu.sheet_view.rightToLeft = True
        _header(ws_ccu, ["مركز الكلفة (نص هبة)", "التاريخ", "رقم القيد", "البيان", "مدين", "دائن"])
        for e in cost_centers["unresolved"]:
            ws_ccu.append([e.get("cost_center", ""), e.get("raw_date", ""), e.get("voucher_no", ""),
                           e.get("narration", ""), float(e.get("debit", 0)), float(e.get("credit", 0))])
            ws_ccu.cell(row=ws_ccu.max_row, column=1).fill = UNRESOLVED_FILL
        for i, w in enumerate([26, 12, 12, 45, 14, 14], start=1):
            ws_ccu.column_dimensions[get_column_letter(i)].width = w

    # ---------------- مطابقة كل صيدلية داخلياً (بون = مرتجع وهمي) ----------------
    # فحص ثانوي (لا يقارن مع هبة إطلاقاً) — يُبنى بعد المطابقة الفعلية أعلاه
    # كي لا يُقرأ خطأً على أنه "النتيجة"، بناءً على طلب المستخدم الصريح 2026-09-30.
    ws0.append(["", ""])
    ws0.append(["مطابقة كل صيدلية داخلياً (بون = مرتجع وهمي) — فحص ثانوي، بلا مقارنة مع هبة", ""])
    ws0.append(["عدد الصيدليات", s["pharmacies_count"]])
    ws0.append(["صيدليات متطابقة داخلياً", s["matched_count"]])
    ws0.append(["صيدليات غير متطابقة داخلياً", s["mismatched_count"]])
    ws0.append(["إجمالي فرق الصيدليات غير المتطابقة داخلياً", float(s["mismatched_amount"])])

    ws1 = wb.create_sheet("مطابقة الصيدليات")
    ws1.sheet_view.rightToLeft = True
    _header(ws1, ["رمز الصيدلية", "اسم الصيدلية", "إجمالي البون (مدين)", "إجمالي المرتجع الوهمي (دائن)",
                  "الفرق", "الحالة (حسابنا)", "علامة النظام المصدر"])
    for p in result["pharmacies"]:
        ws1.append([
            p["code"], p["name"], float(p["total_debit"]), float(p["total_credit"]), float(p["difference"]),
            "مطابق" if p["matched"] else "غير مطابق",
            "غير مطابق" if p["source_flagged_mismatch"] else "مطابق",
        ])
        ws1.cell(row=ws1.max_row, column=6).fill = OK_FILL if p["matched"] else BAD_FILL
    for i, w in enumerate([14, 40, 20, 24, 14, 18, 20], start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("تفاصيل الحركات")
    ws2.sheet_view.rightToLeft = True
    _header(ws2, ["رمز الصيدلية", "اسم الصيدلية", "التاريخ", "رقم السند", "البيان", "مدين", "دائن"])
    for p in result["pharmacies"]:
        for r in p["rows"]:
            ws2.append([p["code"], p["name"], r["date"], r["voucher_no"], r["narration"], float(r["debit"]), float(r["credit"])])
    for i, w in enumerate([14, 34, 12, 12, 50, 14, 14], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ---------------- مطابقة الحركات دواك ⇄ هبة (إضافة لاحقة) ----------------
    payments = result.get("payments")
    if payments is not None:
        ps = payments["summary"]
        # تصحيح 2026-09-30: عمودا "الصيدلية" على كل طرف — الطرف دواك من تسطيح
        # حركاته (pharmacy_name)، والطرف هبة من إسناد مركز الكلفة (تم أعلاه
        # عبر compute_cost_center_reconciliation قبل استدعاء match_dawak_hiba_payments،
        # فكل حركة هبة تحمل hiba_pharmacy_name فعلاً هنا) — حتى يكون واضحاً
        # فوراً أي صيدلية يخص كل سطر بلا حاجة للرجوع لشيت آخر.
        pay_headers = ["الاتجاه", "الصيدلية (دواك)", "تاريخ دواك", "رقم سند دواك", "بيان دواك", "المبلغ",
                       "الصيدلية المطابقة (هبة، عبر مركز الكلفة)", "تاريخ هبة", "رقم سند هبة", "بيان هبة",
                       "فرق الأيام", "ملاحظات"]
        dir_label = {
            "dawak_debit_hiba_credit": "مدين دواك ⇄ دائن هبة",
            "dawak_credit_hiba_debit": "دائن دواك ⇄ مدين هبة",
        }

        def _entry_date_s(e):
            if e is None:
                return ""
            return e["entry_date"].isoformat() if e["entry_date"] else (e["raw_date"] or "")

        def _write_pairs(ws, pairs, fill):
            for p in pairs:
                a, b = p["a"], p["b"]
                ws.append([
                    dir_label.get(p["direction"], p["direction"]),
                    a.get("pharmacy_name", "") if a else "",
                    _entry_date_s(a), a["voucher_no"] if a else "", a["narration"] if a else "",
                    float(p["amount"]),
                    (b.get("hiba_pharmacy_name") or "غير مُسنَد") if b else "",
                    _entry_date_s(b), b["voucher_no"] if b else "", b["narration"] if b else "",
                    p["day_diff"] if p["day_diff"] is not None else "",
                    p["note"],
                ])
            for r in range(2, ws.max_row + 1):
                for c in range(1, len(pay_headers) + 1):
                    ws.cell(row=r, column=c).fill = fill
            for i, w in enumerate([20, 26, 12, 12, 40, 14, 26, 12, 12, 40, 10, 46], start=1):
                ws.column_dimensions[get_column_letter(i)].width = w

        ws3 = wb.create_sheet("مطابقة الدفعات دواك⇄هبة")
        ws3.sheet_view.rightToLeft = True
        _header(ws3, pay_headers)
        _write_pairs(ws3, payments["matched"], OK_FILL)

        # فئة مستقلة (وليست "فرق"): موجود بنفس المبلغ لكن بتاريخ مختلف — فحص ثانوي
        # حسب طلب المستخدم، حتى لا تُخفى المعلومة بأنها "غير موجودة" بالخطأ.
        ws4 = wb.create_sheet("موجود بتاريخ مختلف")
        ws4.sheet_view.rightToLeft = True
        _header(ws4, pay_headers)
        _write_pairs(ws4, payments["found_diff_date"], FOUND_DIFF_DATE_FILL)

        ws5 = wb.create_sheet("فرق حقيقي دواك⇄هبة")
        ws5.sheet_view.rightToLeft = True
        _header(ws5, pay_headers)
        _write_pairs(ws5, payments["true_diff"], BAD_FILL)

        # صفوف إضافية بالملخص لهذه المطابقة
        ws0.append(["", ""])
        ws0.append(["مطابقة الحركات دواك ⇄ هبة", ""])
        ws0.append(["عدد حركات دواك", ps["dawak_entries_count"]])
        ws0.append(["عدد حركات هبة", ps["hiba_entries_count"]])
        ws0.append(["حركات متطابقة", ps["matched_count"]])
        ws0.append(["   منها: نفس التاريخ", ps["matched_same_day"]])
        ws0.append(["   منها: بفارق تاريخ", ps["matched_date_diff"]])
        ws0.append(["موجود بنفس المبلغ بتاريخ مختلف (ليست فرقاً)", ps["found_diff_date_count"]])
        ws0.append(["   إجمالي مبلغها", float(ps["found_diff_date_amount"])])
        ws0.append(["فرق حقيقي (لا يوجد المبلغ إطلاقاً عند الطرف الآخر)", ps["true_diff_count"]])
        ws0.append(["   إجمالي مبلغها", float(ps["true_diff_amount"])])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
