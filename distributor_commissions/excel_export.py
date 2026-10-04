import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from distributor_commissions.engine import ROLE_ORDER, SITUATION_ALONE, SITUATION_GROUP, SITUATION_PAIR

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FLAG_FILL = PatternFill("solid", fgColor="FDE68A")
SUBTOTAL_FILL = PatternFill("solid", fgColor="E5E7EB")
# تصحيح 2026-09-28 (طلب المستخدم الصريح، ألوان مميّزة لحالتين جديدتين):
#  1) FUZZY_FILL (برتقالي مميَّز، مختلف عن FLAG_FILL الأصفر الفاتح العام):
#     مشارك استُخرج اسمه من نص البيان بمطابقة تقريبية فقط (كنية/خطأ إملائي)
#     — يحتاج تأكيداً يدوياً من المستخدم رغم احتسابه طبيعياً بالعمولة.
#  2) UNRESOLVED_FILL (أحمر — طُلب صراحة "باللون الأحمر"): حركة لا يوجد لها
#     أي مشارك معروف إطلاقاً، لا بعمود خام ولا حتى بمطابقة تقريبية من البيان.
FUZZY_FILL = PatternFill("solid", fgColor="FDBA74")
UNRESOLVED_FILL = PatternFill("solid", fgColor="FCA5A5")
UNRESOLVED_FONT = Font(color="7F1D1D", bold=True)

TIER_LABEL = {"car": "سيارة", "moto": "دراجة"}
SITUATION_ORDER = [SITUATION_ALONE, SITUATION_PAIR, SITUATION_GROUP]


def _header(ws, headers):
    ws.append(headers)
    r = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = f"A{r + 1}"


def build_workbook(parsed, result, summary, meta, situation_breakdown=None, return_situation_breakdown=None,
                    role_breakdown=None, return_role_breakdown=None):
    wb = Workbook()

    ws0 = wb.active
    ws0.title = "الملخص"
    ws0.sheet_view.rightToLeft = True
    ws0.append(["ملخص عمولة تحصيل الموزعين الداخليين — نقدية + مرتجعات"])
    ws0.append([])
    ws0.append(["ملف الحركة", meta.get("ledger_file_name", "")])
    ws0.append(["عدد حركات التحصيل النقدي المحتسبة", summary["rows_count"]])
    ws0.append(["عدد حركات المرتجع المحتسبة", summary["return_rows_count"]])
    ws0.append(["عدد الموزعين", summary["distributors_count"]])
    ws0.append(["إجمالي عمولة التحصيل (قبل المرتجعات)", float(summary["collection_total"])])
    ws0.append(["إجمالي خصم المرتجعات", float(summary["return_total"])])
    ws0.append(["الصافي النهائي", float(summary["total_commission"])])
    ws0.append([])
    ws0.append(["حركات مستبعدة كلياً (بلا عمولة):"])
    ws0.append(["  — حركات حساب مستودع (عمود الحساب المقابل)", parsed.warehouse_excluded_count,
                float(parsed.warehouse_excluded_amount)])
    for marker, amt in sorted(parsed.excluded_marker_totals.items()):
        ws0.append([f"  — علامة غير-شخصية: {marker}", "", float(amt)])
    ws0.append([])
    if parsed.flagged_rows:
        ws0.append([f"⚠ يوجد {len(parsed.flagged_rows)} حركة تحتاج مراجعة يدوية — راجع شيت 'حركات للمراجعة'."])
    if parsed.fuzzy_biyad_rows:
        ws0.append([f"⚠ يوجد {len(parsed.fuzzy_biyad_rows)} حركة فيها مشارك استُخرج اسمه من نص البيان بمطابقة "
                    f"تقريبية فقط (لوّنت برتقالياً) — راجع شيت 'مطابقة تقريبية من البيان' للتأكد يدوياً."])
    if parsed.unresolved_collection_rows or parsed.unresolved_return_rows:
        total_unresolved = len(parsed.unresolved_collection_rows) + len(parsed.unresolved_return_rows)
        ws0.append([f"⚠ يوجد {total_unresolved} حركة بلا أي موزع معروف إطلاقاً (لا عمود خام ولا حتى مطابقة "
                    f"تقريبية بالبيان، لوّنت أحمر) — راجع شيت 'حركات بلا أي موزع معروف' واحتسبها يدوياً."])
    ws0.append(["ملاحظة: خصم المرتجعات نسبة ثابتة 0.5% لكل مشارك (وليس النسب المتدرجة"
                " المستخدمة في التحصيل النقدي) — بحسب مطابقة شيت 'تقرير نتيجة' المرجعي"
                " الذي زوّدنا به المستخدم للتحقق؛ راجع توثيق الكود لتفاصيل هذا القرار."])
    ws0.column_dimensions["A"].width = 46
    ws0.column_dimensions["B"].width = 16
    ws0.column_dimensions["C"].width = 16

    ws1 = wb.create_sheet("عمولة الموزعين (الصافي)")
    ws1.sheet_view.rightToLeft = True
    _header(ws1, ["اسم الموزع", "عمولة التحصيل النقدي", "خصم المرتجعات", "الصافي"])
    collection = result["collection_commission"]
    returns = result["return_commission"]
    for name in sorted(result["totals"]):
        ws1.append([
            name,
            float(collection.get(name, 0)),
            float(returns.get(name, 0)),
            float(result["totals"][name]),
        ])
    for i, w in enumerate([30, 20, 18, 16], start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("تفصيل حركات التحصيل")
    ws2.sheet_view.rightToLeft = True
    _header(ws2, ["البيان", "المبلغ (مدين)", "عدد المشاركين", "المشاركون (اسم/فئة/مصدر)"])
    for row in parsed.collection_rows:
        parts_str = "، ".join(
            f"{p.name} ({TIER_LABEL.get(p.tier, p.tier)}"
            f"{'، من البيان' if p.role_col is None else ''}"
            f"{'، مطابقة تقريبية ⚠' if p.match_kind == 'fuzzy' else ''})"
            for p in row.participants
        )
        ws2.append([row.bayan, float(row.debit), len(row.participants), parts_str])
        if any(p.match_kind == "fuzzy" for p in row.participants):
            for c in range(1, 5):
                ws2.cell(row=ws2.max_row, column=c).fill = FUZZY_FILL
    for i, w in enumerate([70, 16, 14, 60], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws3 = wb.create_sheet("تفصيل المرتجعات")
    ws3.sheet_view.rightToLeft = True
    _header(ws3, ["البيان", "المبلغ", "عدد المشاركين", "المشاركون"])
    for row in parsed.return_rows:
        parts_str = "، ".join(p.name for p in row.participants)
        ws3.append([row.bayan, float(row.amount), len(row.participants), parts_str])
    for i, w in enumerate([40, 16, 14, 60], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------
    # تفصيل كل موزع حسب الحالة (منفرد / مع شخص آخر / مجموعة) — الطلب الجديد:
    # "تفصيل كل موزع منفرداً مع الربح المالي لكل حالة والعمولة لكل حالة".
    # سطر واحد لكل (موزع × حالة)؛ المجموع عبر الحالات الثلاث لكل موزع يساوي
    # عمولته الإجمالية في شيت "عمولة الموزعين (الصافي)" تماماً (لا تُغيَّر
    # النسب هنا إطلاقاً — فقط إعادة عرض لنفس الاحتساب بحسب الحالة).
    # ------------------------------------------------------------------
    ws5 = wb.create_sheet("تفصيل كل موزع حسب الحالة")
    ws5.sheet_view.rightToLeft = True
    _header(ws5, ["اسم الموزع", "الحالة", "عدد الحركات", "إجمالي المبلغ (مدين)", "العمولة"])
    situation_breakdown = situation_breakdown or {}
    for name in sorted(result["totals"]):
        situations = situation_breakdown.get(name, {})
        subtotal_count = 0
        subtotal_amount = 0.0
        subtotal_commission = 0.0
        for sit in SITUATION_ORDER:
            b = situations.get(sit, {"count": 0, "amount": 0, "commission": 0})
            ws5.append([name, sit, b["count"], float(b["amount"]), float(b["commission"])])
            subtotal_count += b["count"]
            subtotal_amount += float(b["amount"])
            subtotal_commission += float(b["commission"])
        r = ws5.max_row + 1
        ws5.append([name, "إجمالي الموزع (كل الحالات)", subtotal_count, subtotal_amount, subtotal_commission])
        for c in range(1, 6):
            ws5.cell(row=r, column=c).fill = SUBTOTAL_FILL
            ws5.cell(row=r, column=c).font = Font(bold=True)
    if not result["totals"]:
        ws5.append(["لا توجد بيانات."])
    for i, w in enumerate([30, 26, 14, 20, 16], start=1):
        ws5.column_dimensions[get_column_letter(i)].width = w

    ws5b = wb.create_sheet("تفصيل المرتجعات حسب الحالة")
    ws5b.sheet_view.rightToLeft = True
    ws5b.append(["ملاحظة: نسبة خصم المرتجع ثابتة 0.5% لكل مشارك بصرف النظر عن الحالة"
                 " (منفرد/مع غيره) — هذا الشيت يوضّح فقط توزّع عدد/مبلغ/خصم المرتجعات"
                 " على كل حالة لكل موزع للشفافية، وليس نسباً مختلفة حسب الحالة."])
    ws5b.append([])
    _header(ws5b, ["اسم الموزع", "الحالة", "عدد حركات المرتجع", "إجمالي مبلغ المرتجع", "خصم العمولة"])
    return_situation_breakdown = return_situation_breakdown or {}
    return_names = sorted(set(result["return_commission"]) | set(return_situation_breakdown))
    for name in return_names:
        situations = return_situation_breakdown.get(name, {})
        subtotal_count = 0
        subtotal_amount = 0.0
        subtotal_commission = 0.0
        for sit in SITUATION_ORDER:
            b = situations.get(sit, {"count": 0, "amount": 0, "commission": 0})
            ws5b.append([name, sit, b["count"], float(b["amount"]), float(b["commission"])])
            subtotal_count += b["count"]
            subtotal_amount += float(b["amount"])
            subtotal_commission += float(b["commission"])
        r = ws5b.max_row + 1
        ws5b.append([name, "إجمالي الموزع (كل الحالات)", subtotal_count, subtotal_amount, subtotal_commission])
        for c in range(1, 6):
            ws5b.cell(row=r, column=c).fill = SUBTOTAL_FILL
            ws5b.cell(row=r, column=c).font = Font(bold=True)
    if not return_names:
        ws5b.append(["لا توجد بيانات مرتجعات."])
    for i, w in enumerate([30, 26, 16, 20, 16], start=1):
        ws5b.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------
    # تفصيل كل موزع حسب دوره/حالته الفعلية (سائق السيارة/مساعد/سائق الدراجة/
    # دراجة 2/استثناء/من البيان) — طلب المستخدم الصريح 2026-09-28. بُعد
    # مختلف عن "الحالة" أعلاه (عدد المشاركين) — راجع توثيق
    # compute_collection_role_breakdown في engine.py. عمود "النسبة الفعلية"
    # = العمولة ÷ المبلغ لكل حزمة (موزع×دور)، مُشتقّة من الأرقام الحقيقية.
    # ------------------------------------------------------------------
    ws6 = wb.create_sheet("تفصيل كل موزع حسب الدور")
    ws6.sheet_view.rightToLeft = True
    ws6.append(["\"الدور\" هنا = العمود الذي ظهر فيه الموزع فعلياً بكل حركة (سائق السيارة/مساعد/سائق الدراجة/"
                "دراجة 2/استثناء)، أو \"من نص البيان\" حين استُخرج اسمه من آلية \"بيد السيد\" بلا عمود خام. "
                "النسبة الفعلية = العمولة ÷ المبلغ لهذا الدور تحديداً — قد تختلف بين حركات نفس الدور لنفس "
                "الموزع إن اختلف عدد المشاركين بالحركة (النسبة الرسمية تعتمد على العدد أيضاً، لا الدور وحده)."])
    ws6.append([])
    _header(ws6, ["اسم الموزع", "الدور/الحالة", "عدد الحركات", "إجمالي المبلغ (مدين)", "العمولة", "النسبة الفعلية"])
    role_breakdown = role_breakdown or {}
    for name in sorted(result["totals"]):
        roles = role_breakdown.get(name, {})
        subtotal_count = 0
        subtotal_amount = 0.0
        subtotal_commission = 0.0
        for role in ROLE_ORDER:
            b = roles.get(role)
            if not b or b["count"] == 0:
                continue
            amount_v, commission_v = float(b["amount"]), float(b["commission"])
            rate_cell_value = (commission_v / amount_v) if amount_v else None
            r = ws6.max_row + 1
            ws6.append([name, role, b["count"], amount_v, commission_v,
                        rate_cell_value if rate_cell_value is not None else ""])
            if rate_cell_value is not None:
                ws6.cell(row=r, column=6).number_format = "0.00%"
            subtotal_count += b["count"]
            subtotal_amount += amount_v
            subtotal_commission += commission_v
        r = ws6.max_row + 1
        ws6.append([name, "إجمالي الموزع (كل الأدوار)", subtotal_count, subtotal_amount, subtotal_commission,
                    (subtotal_commission / subtotal_amount) if subtotal_amount else ""])
        if subtotal_amount:
            ws6.cell(row=r, column=6).number_format = "0.00%"
        for c in range(1, 7):
            ws6.cell(row=r, column=c).fill = SUBTOTAL_FILL
            ws6.cell(row=r, column=c).font = Font(bold=True)
    if not result["totals"]:
        ws6.append(["لا توجد بيانات."])
    for i, w in enumerate([30, 34, 14, 20, 16, 16], start=1):
        ws6.column_dimensions[get_column_letter(i)].width = w

    ws6b = wb.create_sheet("تفصيل المرتجعات حسب الدور")
    ws6b.sheet_view.rightToLeft = True
    ws6b.append(["ملاحظة: نسبة خصم المرتجع ثابتة 0.5% لكل مشارك بصرف النظر عن الدور — هذا الشيت يوضّح فقط "
                 "توزّع عدد/مبلغ/خصم المرتجعات على كل دور لكل موزع للشفافية."])
    ws6b.append([])
    _header(ws6b, ["اسم الموزع", "الدور/الحالة", "عدد حركات المرتجع", "إجمالي مبلغ المرتجع", "خصم العمولة"])
    return_role_breakdown = return_role_breakdown or {}
    return_role_names = sorted(set(result["return_commission"]) | set(return_role_breakdown))
    for name in return_role_names:
        roles = return_role_breakdown.get(name, {})
        subtotal_count = 0
        subtotal_amount = 0.0
        subtotal_commission = 0.0
        for role in ROLE_ORDER:
            b = roles.get(role)
            if not b or b["count"] == 0:
                continue
            ws6b.append([name, role, b["count"], float(b["amount"]), float(b["commission"])])
            subtotal_count += b["count"]
            subtotal_amount += float(b["amount"])
            subtotal_commission += float(b["commission"])
        r = ws6b.max_row + 1
        ws6b.append([name, "إجمالي الموزع (كل الأدوار)", subtotal_count, subtotal_amount, subtotal_commission])
        for c in range(1, 6):
            ws6b.cell(row=r, column=c).fill = SUBTOTAL_FILL
            ws6b.cell(row=r, column=c).font = Font(bold=True)
    if not return_role_names:
        ws6b.append(["لا توجد بيانات مرتجعات."])
    for i, w in enumerate([30, 34, 16, 20, 16], start=1):
        ws6b.column_dimensions[get_column_letter(i)].width = w

    ws4 = wb.create_sheet("حركات للمراجعة")
    ws4.sheet_view.rightToLeft = True
    _header(ws4, ["السبب", "البيان", "المبلغ", "الاسم/الأسماء"])
    for f in parsed.flagged_rows:
        r = ws4.max_row + 1
        ws4.append([f["reason"], f["bayan"], float(f["debit"]), f["name"]])
        for c in range(1, 5):
            ws4.cell(row=r, column=c).fill = FLAG_FILL
    if not parsed.flagged_rows:
        ws4.append(["لا توجد حركات مُعلَّمة للمراجعة في هذا التشغيل."])
    for i, w in enumerate([30, 60, 14, 30], start=1):
        ws4.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------
    # الحالة الأولى (طلب المستخدم 2026-09-28): مشاركون استُخرجت أسماؤهم من
    # نص البيان بمطابقة تقريبية فقط (كنية/خطأ إملائي طفيف) — محتسَبون
    # بالعمولة طبيعياً (راجع شيت "تفصيل حركات التحصيل" الملوَّن برتقالي)،
    # لكن مُجمَّعون هنا أيضاً لمراجعة سريعة مخصصة.
    # ------------------------------------------------------------------
    ws7 = wb.create_sheet("مطابقة تقريبية من البيان")
    ws7.sheet_view.rightToLeft = True
    ws7.append(["هذه حركات فيها مشارك واحد أو أكثر استُخرج اسمه من نص البيان (\"بيد السيد ...\") بمطابقة "
                "تقريبية فقط — اسم جزئي (كنية) أو تقارب إملائي (احتمال اختلاف بحرف أو أكثر عن الاسم الفعلي)، "
                "وليس تطابقاً حرفياً أو اسماً مؤكَّداً يدوياً مسبقاً. عمولتها محتسَبة طبيعياً بافتراض صحة "
                "المطابقة — راجعها للتأكد."])
    _header(ws7, ["البيان", "المبلغ (مدين)", "الاسم/الأسماء بمطابقة تقريبية", "كل المشاركين بالحركة"])
    for f in parsed.fuzzy_biyad_rows:
        r = ws7.max_row + 1
        ws7.append([f["bayan"], float(f["debit"]), f["names"], f["all_names"]])
        for c in range(1, 5):
            ws7.cell(row=r, column=c).fill = FUZZY_FILL
    if not parsed.fuzzy_biyad_rows:
        ws7.append(["لا توجد حركات بمطابقة تقريبية من البيان في هذا التشغيل."])
    for i, w in enumerate([70, 16, 40, 50], start=1):
        ws7.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------
    # الحالة الثانية (طلب المستخدم 2026-09-28): حركات بلا أي موزع معروف
    # إطلاقاً — لا عمود DIST_COLS خام ولا حتى مطابقة تقريبية من البيان.
    # سابقاً كانت تُحذف صامتة بالكامل؛ الآن تبقى ظاهرة هنا (أحمر) بدل
    # الاختفاء، بلا أي عمولة محتسَبة لها (لا يوجد من يُنسب له المبلغ).
    # ------------------------------------------------------------------
    ws8 = wb.create_sheet("حركات بلا أي موزع معروف")
    ws8.sheet_view.rightToLeft = True
    ws8.append(["هذه حركات (تحصيل نقدي أو مرتجع) لا يوجد لها أي مشارك معروف إطلاقاً — لا بعمود موزع خام ولا "
                "حتى بمطابقة تقريبية من نص البيان. لم تُحتسب لأي موزع (بلا اسم يُنسب له المبلغ) ولم تُحذف "
                "صامتة كما كان يحدث سابقاً — احتسبها يدوياً بعد تحديد الموزع الفعلي."])
    _header(ws8, ["المصدر", "البيان", "المبلغ"])
    for f in parsed.unresolved_collection_rows:
        r = ws8.max_row + 1
        ws8.append(["تحصيل نقدي (دفتر الأستاذ)", f["bayan"], float(f["debit"])])
        for c in range(1, 4):
            ws8.cell(row=r, column=c).fill = UNRESOLVED_FILL
            ws8.cell(row=r, column=c).font = UNRESOLVED_FONT
    for f in parsed.unresolved_return_rows:
        r = ws8.max_row + 1
        ws8.append(["مرتجع", f["bayan"], float(f["debit"])])
        for c in range(1, 4):
            ws8.cell(row=r, column=c).fill = UNRESOLVED_FILL
            ws8.cell(row=r, column=c).font = UNRESOLVED_FONT
    if not parsed.unresolved_collection_rows and not parsed.unresolved_return_rows:
        ws8.append(["لا توجد حركات بلا موزع معروف في هذا التشغيل."])
    for i, w in enumerate([26, 70, 16], start=1):
        ws8.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
