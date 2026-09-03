from decimal import Decimal

import openpyxl
from django.contrib import messages
from django.core.files.base import ContentFile
from django.http import FileResponse, Http404
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from core.audit import log_action
from core.decorators import module_required
from core.models import AuditLog
from reconciliation_hs.engine import LedgerParseError, parse_ledger, reconcile
from reconciliation_hs.excel_export import build_workbook
from reconciliation_hs.forms import ReconciliationUploadForm
from reconciliation_hs.models import ReconciliationRun

PREVIEW_LIMIT = 40


@module_required("reconciliation_hs")
def index(request):
    form = ReconciliationUploadForm()
    runs = ReconciliationRun.objects.select_related("created_by")[:15]
    return render(request, "reconciliation_hs/index.html", {"form": form, "runs": runs})


@module_required("reconciliation_hs", edit=True)
def run_reconciliation(request):
    if request.method != "POST":
        return redirect("reconciliation_hs:index")

    form = ReconciliationUploadForm(request.POST, request.FILES)
    if not form.is_valid():
        runs = ReconciliationRun.objects.select_related("created_by")[:15]
        return render(request, "reconciliation_hs/index.html", {"form": form, "runs": runs})

    sadana_file = form.cleaned_data["sadana_file"]
    hiba_file = form.cleaned_data["hiba_file"]

    try:
        hiba_entries = parse_ledger(hiba_file, "hiba")
        sadana_entries = parse_ledger(sadana_file, "sadana")
    except LedgerParseError as exc:
        messages.error(request, f"تعذّرت قراءة أحد الملفين: {exc}")
        log_action(request, AuditLog.Action.ERROR, str(exc), module_code="reconciliation_hs")
        return redirect("reconciliation_hs:index")
    except Exception as exc:  # noqa: BLE001
        messages.error(request, f"حدث خطأ غير متوقع أثناء قراءة الملفات: {exc}")
        log_action(request, AuditLog.Action.ERROR, str(exc), module_code="reconciliation_hs")
        return redirect("reconciliation_hs:index")

    if not hiba_entries or not sadana_entries:
        messages.error(request, "لم يتم العثور على أي حركات صالحة داخل أحد الملفين. الرجاء التأكد من تنسيق الملف.")
        return redirect("reconciliation_hs:index")

    result = reconcile(hiba_entries, sadana_entries)
    summary = result["summary"]

    run = ReconciliationRun.objects.create(
        created_by=request.user,
        hiba_file_name=hiba_file.name,
        sadana_file_name=sadana_file.name,
        hiba_count=summary["hiba_count"],
        sadana_count=summary["sadana_count"],
        matched_count=summary["matched_count"],
        matched_same_day=summary["matched_same_day"],
        matched_date_diff=summary["matched_date_diff"],
        hiba_only_count=summary["hiba_only_count"],
        sadana_only_count=summary["sadana_only_count"],
        hiba_only_amount=Decimal(summary["hiba_only_amount"]),
        sadana_only_amount=Decimal(summary["sadana_only_amount"]),
        found_diff_date_count=summary["found_diff_date_count"],
        found_diff_date_amount=Decimal(summary["found_diff_date_amount"]),
        hiba_period_start=summary["hiba_period"][0],
        hiba_period_end=summary["hiba_period"][1],
        sadana_period_start=summary["sadana_period"][0],
        sadana_period_end=summary["sadana_period"][1],
        has_errors=summary["has_errors"],
    )

    workbook_buf = build_workbook(result, {
        "run_time": timezone.localtime(run.created_at).strftime("%Y-%m-%d %H:%M"),
        "hiba_file_name": hiba_file.name,
        "sadana_file_name": sadana_file.name,
    })
    run.result_file.save(f"مطابقة_هبة_سدانة_{run.pk}.xlsx", ContentFile(workbook_buf.read()), save=True)

    log_action(
        request, AuditLog.Action.RUN,
        f"مطابقة هبة/سدانة #{run.pk}: {summary['matched_count']} متطابقة، "
        f"{summary['hiba_only_count']} فقط بهبة، {summary['sadana_only_count']} فقط بسدانة",
        module_code="reconciliation_hs",
        meta={"run_id": run.pk},
    )

    msgs = []
    if summary["has_errors"]:
        msgs.append("توجد فروقات تحتاج مراجعتك — راجع التفاصيل أدناه.")
    if summary["found_diff_date_count"]:
        msgs.append(
            f"تم العثور على {summary['found_diff_date_count']} حركة بنفس المبلغ عند الطرف الآخر لكن بتاريخ مختلف "
            f"(إجمالي {summary['found_diff_date_amount']}) — لا تُحسب فرقاً حقيقياً، لكن يُستحسن التحقق اليدوي منها.",
        )
    if msgs:
        messages.warning(request, " ".join(msgs))
    else:
        messages.success(request, "تمت المطابقة بنجاح ولم يتم العثور على أي أخطاء.")

    request.session[f"reconc_preview_{run.pk}"] = _build_preview(result)
    return redirect("reconciliation_hs:result", pk=run.pk)


def _build_preview(result):
    def pair_row(p, only_side):
        e = p.a if only_side == "hiba" else p.b
        kind = "مدين" if e.debit > 0 else "دائن"
        return {
            "date": e.entry_date.isoformat() if e.entry_date else e.raw_date,
            "voucher": e.voucher_no,
            "narration": e.narration,
            "kind": kind,
            "amount": str(p.amount),
            "note": p.note,
        }

    def matched_row(p):
        h, s = p.a, p.b
        return {
            "h_date": h.entry_date.isoformat() if h.entry_date else h.raw_date,
            "h_voucher": h.voucher_no,
            "h_narration": h.narration,
            "s_date": s.entry_date.isoformat() if s.entry_date else s.raw_date,
            "s_voucher": s.voucher_no,
            "s_narration": s.narration,
            "amount": str(p.amount),
            "day_diff": p.day_diff,
            "note": p.note,
            "status": p.status,
        }

    def found_diff_date_row(p):
        # a = حركة هبة دائماً، b = حركة سدانة دائماً (انظر _unmatched_pair)
        h, s = p.a, p.b
        return {
            "amount": str(p.amount),
            "hiba_date": h.entry_date.isoformat() if h.entry_date else (h.raw_date or ""),
            "hiba_voucher": h.voucher_no,
            "sadana_date": s.entry_date.isoformat() if s.entry_date else (s.raw_date or ""),
            "sadana_voucher": s.voucher_no,
            "note": p.note,
        }

    return {
        "hiba_only": [pair_row(p, "hiba") for p in result["hiba_only"][:PREVIEW_LIMIT]],
        "sadana_only": [pair_row(p, "sadana") for p in result["sadana_only"][:PREVIEW_LIMIT]],
        "matched_diff": [matched_row(p) for p in result["matched"] if p.status == "matched_date_diff"][:PREVIEW_LIMIT],
        "found_diff_date": [found_diff_date_row(p) for p in result["found_diff_date"][:PREVIEW_LIMIT]],
    }


@module_required("reconciliation_hs")
def result(request, pk):
    run = get_object_or_404(ReconciliationRun, pk=pk)
    preview = request.session.get(f"reconc_preview_{run.pk}")
    if preview is None:
        preview = _reload_preview_from_file(run)
    return render(request, "reconciliation_hs/result.html", {"run": run, "preview": preview})


def _reload_preview_from_file(run):
    """في حال لم تعد المعاينة محفوظة في الجلسة (مثلاً بعد إعادة تشغيل الخادم)، نعيد قراءتها من ملف النتيجة."""
    if not run.result_file:
        return {"hiba_only": [], "sadana_only": [], "matched_diff": []}
    try:
        wb = openpyxl.load_workbook(run.result_file.path, data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        return {"hiba_only": [], "sadana_only": [], "matched_diff": []}

    def read_only_sheet(name):
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))[1:PREVIEW_LIMIT + 1]
        out = []
        for r in rows:
            out.append({
                "date": r[0], "voucher": r[1], "narration": r[2], "kind": r[3],
                "amount": r[4], "note": r[6],
            })
        return out

    def read_matched_diff():
        ws = wb["الحركات المتطابقة"]
        rows = list(ws.iter_rows(values_only=True))[1:]
        out = []
        for r in rows:
            if r[11] not in ("", None, 0):
                out.append({
                    "h_date": r[1], "h_voucher": r[2], "h_narration": r[3],
                    "s_date": r[6], "s_voucher": r[7], "s_narration": r[8],
                    "amount": r[4], "day_diff": r[11], "note": r[12], "status": "matched_date_diff",
                })
            if len(out) >= PREVIEW_LIMIT:
                break
        return out

    def read_found_diff_date():
        ws = wb["موجود بتاريخ مختلف"]
        rows = list(ws.iter_rows(values_only=True))[1:PREVIEW_LIMIT + 1]
        out = []
        for r in rows:
            out.append({
                "hiba_date": r[1], "hiba_voucher": r[2],
                "sadana_date": r[6], "sadana_voucher": r[7],
                "amount": r[4], "note": r[12],
            })
        return out

    preview = {
        "hiba_only": read_only_sheet("موجود في هبة فقط"),
        "sadana_only": read_only_sheet("موجود في سدانة فقط"),
        "matched_diff": read_matched_diff(),
        "found_diff_date": read_found_diff_date(),
    }
    wb.close()
    return preview


@module_required("reconciliation_hs")
def download(request, pk):
    run = get_object_or_404(ReconciliationRun, pk=pk)
    if not run.result_file:
        raise Http404("لا يوجد ملف نتيجة لهذه العملية.")
    log_action(request, AuditLog.Action.EXPORT, f"تنزيل نتيجة مطابقة #{run.pk}", module_code="reconciliation_hs")
    return FileResponse(run.result_file.open("rb"), as_attachment=True,
                         filename=f"مطابقة_هبة_سدانة_{run.pk}.xlsx")
