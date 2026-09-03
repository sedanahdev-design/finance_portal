"""بناء ملف الإكسل الناتج عن عملية المطابقة، بتنسيق واضح ومرتّب."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
MATCH_FILL = PatternFill("solid", fgColor="E7F7EF")
DIFF_DATE_FILL = PatternFill("solid", fgColor="FFF7E0")
ONLY_FILL = PatternFill("solid", fgColor="FDEAEA")
FOUND_DIFF_DATE_FILL = PatternFill("solid", fgColor="FDE9C8")  # برتقالي: موجود بمبلغ مطابق لكن بتاريخ مختلف
THIN = Side(style="thin", color="D9DCE3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = f"A{row + 1}"


def _autofit(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _sheet_rtl(ws):
    ws.sheet_view.rightToLeft = True


def build_workbook(result, run_meta):
    wb = Workbook()

    # ---------------- ملخص ----------------
    ws = wb.active
    ws.title = "الملخص"
    _sheet_rtl(ws)
    summary = result["summary"]
    rows = [
        ["مطابقة كشفي حساب هبة - سدانة", ""],
        ["تاريخ التشغيل", run_meta.get("run_time", "")],
        ["ملف هبة", run_meta.get("hiba_file_name", "")],
        ["ملف سدانة", run_meta.get("sadana_file_name", "")],
        ["", ""],
        ["عدد حركات هبة", summary["hiba_count"]],
        ["عدد حركات سدانة", summary["sadana_count"]],
        ["عدد الحركات المتطابقة", summary["matched_count"]],
        ["   منها: نفس التاريخ تماماً", summary["matched_same_day"]],
        ["   منها: بفارق تاريخ", summary["matched_date_diff"]],
        ["موجود بنفس المبلغ بتاريخ مختلف (فحص ثانوي، ليست فرقاً)", summary["found_diff_date_count"]],
        ["   إجمالي مبلغها", float(summary["found_diff_date_amount"])],
        ["حركات موجودة في هبة فقط (بلا مقابل)", summary["hiba_only_count"]],
        ["   إجمالي مبلغها", float(summary["hiba_only_amount"])],
        ["حركات موجودة في سدانة فقط (بلا مقابل)", summary["sadana_only_count"]],
        ["   إجمالي مبلغها", float(summary["sadana_only_amount"])],
        ["", ""],
        ["النتيجة النهائية",
         "⚠ توجد فروقات تحتاج مراجعة" if summary["has_errors"] else "✔ مطابقة تامة، لا توجد أخطاء"],
    ]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=1, max_row=len(rows), max_col=2):
        row[0].font = Font(bold=True, size=11)
    ws["A1"].font = Font(bold=True, size=14)
    ws.cell(row=len(rows), column=2).font = Font(
        bold=True, size=12,
        color="C0392B" if summary["has_errors"] else "1E8449",
    )
    _autofit(ws, [38, 42])

    # ---------------- المطابقة ----------------
    ws2 = wb.create_sheet("الحركات المتطابقة")
    _sheet_rtl(ws2)
    headers = [
        "الاتجاه", "تاريخ هبة", "رقم سند هبة", "بيان هبة", "مبلغ هبة", "نوع الحركة (هبة)",
        "تاريخ سدانة", "رقم سند سدانة", "بيان سدانة", "مبلغ سدانة", "نوع الحركة (سدانة)",
        "فرق الأيام", "ملاحظات",
    ]
    ws2.append(headers)
    _style_header(ws2, len(headers))
    dir_label = {
        "hiba_debit_sadana_credit": "مدين هبة ⇄ دائن سدانة",
        "hiba_credit_sadana_debit": "دائن هبة ⇄ مدين سدانة",
    }
    for p in result["matched"]:
        h, s = p.a, p.b
        h_kind = "مدين" if h.debit > 0 else "دائن"
        s_kind = "مدين" if s.debit > 0 else "دائن"
        ws2.append([
            dir_label.get(p.direction, p.direction),
            h.entry_date.isoformat() if h.entry_date else h.raw_date,
            h.voucher_no, h.narration, float(p.amount), h_kind,
            s.entry_date.isoformat() if s.entry_date else s.raw_date,
            s.voucher_no, s.narration, float(p.amount), s_kind,
            p.day_diff if p.day_diff is not None else "",
            p.note,
        ])
    for r in range(2, ws2.max_row + 1):
        fill = MATCH_FILL if ws2.cell(row=r, column=12).value in ("", 0) else DIFF_DATE_FILL
        for c in range(1, len(headers) + 1):
            ws2.cell(row=r, column=c).border = BORDER
            ws2.cell(row=r, column=c).fill = fill
    _autofit(ws2, [20, 12, 12, 42, 14, 12, 12, 12, 42, 14, 12, 10, 34])

    # ---------------- موجود بتاريخ مختلف (فحص ثانوي) ----------------
    # حركات لم تُطابَق تلقائياً ضمن مجموعة مبلغها (نقص عدد الحركات بنفس المبلغ عند
    # الطرف الآخر)، لكن الفحص الثانوي وجد نفس المبلغ تماماً عند الطرف الآخر بتاريخ
    # مختلف. لا تُحتسب "فرق" — فئة مستقلة تحتاج مراجعة يدوية بدل تجاهلها كـ"غير موجود".
    ws2b = wb.create_sheet("موجود بتاريخ مختلف")
    _sheet_rtl(ws2b)
    ws2b.append(headers)
    _style_header(ws2b, len(headers))
    for p in result["found_diff_date"]:
        h, s = p.a, p.b
        h_kind = "مدين" if h.debit > 0 else "دائن"
        s_kind = "مدين" if s.debit > 0 else "دائن"
        ws2b.append([
            dir_label.get(p.direction, p.direction),
            h.entry_date.isoformat() if h.entry_date else h.raw_date,
            h.voucher_no, h.narration, float(p.amount), h_kind,
            s.entry_date.isoformat() if s.entry_date else s.raw_date,
            s.voucher_no, s.narration, float(p.amount), s_kind,
            p.day_diff if p.day_diff is not None else "",
            p.note,
        ])
    for r in range(2, ws2b.max_row + 1):
        for c in range(1, len(headers) + 1):
            ws2b.cell(row=r, column=c).border = BORDER
            ws2b.cell(row=r, column=c).fill = FOUND_DIFF_DATE_FILL
    _autofit(ws2b, [20, 12, 12, 42, 14, 12, 12, 12, 42, 14, 12, 10, 46])

    # ---------------- هبة فقط ----------------
    ws3 = wb.create_sheet("موجود في هبة فقط")
    _sheet_rtl(ws3)
    headers3 = ["التاريخ", "رقم السند", "البيان", "نوع الحركة", "المبلغ", "الحساب المقابل", "ملاحظات"]
    ws3.append(headers3)
    _style_header(ws3, len(headers3))
    for p in result["hiba_only"]:
        e = p.a
        kind = "مدين" if e.debit > 0 else "دائن"
        ws3.append([
            e.entry_date.isoformat() if e.entry_date else e.raw_date,
            e.voucher_no, e.narration, kind, float(p.amount), e.contra_account, p.note,
        ])
    for r in range(2, ws3.max_row + 1):
        for c in range(1, len(headers3) + 1):
            ws3.cell(row=r, column=c).border = BORDER
            ws3.cell(row=r, column=c).fill = ONLY_FILL
    _autofit(ws3, [12, 12, 46, 10, 14, 30, 40])

    # ---------------- سدانة فقط ----------------
    ws4 = wb.create_sheet("موجود في سدانة فقط")
    _sheet_rtl(ws4)
    ws4.append(headers3)
    _style_header(ws4, len(headers3))
    for p in result["sadana_only"]:
        e = p.b
        kind = "مدين" if e.debit > 0 else "دائن"
        ws4.append([
            e.entry_date.isoformat() if e.entry_date else e.raw_date,
            e.voucher_no, e.narration, kind, float(p.amount), e.contra_account, p.note,
        ])
    for r in range(2, ws4.max_row + 1):
        for c in range(1, len(headers3) + 1):
            ws4.cell(row=r, column=c).border = BORDER
            ws4.cell(row=r, column=c).fill = ONLY_FILL
    _autofit(ws4, [12, 12, 46, 10, 14, 30, 40])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
