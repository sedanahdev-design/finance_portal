import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def safe_sheet_title(text, fallback="مادة"):
    cleaned = INVALID_SHEET_CHARS.sub(" ", str(text or "")).strip()
    return (cleaned[:31] or fallback)


HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FLAG_FILL = PatternFill("solid", fgColor="FEF3C7")  # مجموعة بلا عرض مفرق صالح / سطر مستبعد — يحتاج مراجعة
IGNORED_FILL = PatternFill("solid", fgColor="F3F4F6")  # مادة مستبعدة بالكامل (بلا عرض مميز إطلاقاً)
TOTAL_FILL = PatternFill("solid", fgColor="FEF08A")  # صف الإجمالي — أصفر، يطابق تلوين ملف المستخدم المرجعي
EXCLUDED_ROW_FILL = PatternFill("solid", fgColor="FEE2E2")  # سطر صافٍ سالب/صفري مستبعد من المجموع

GROUP_HEADERS = ["اسم المادة", "عرض المفرق", "مجموع الكمية", "مجموع الهدايا",
                  "ناتج المعادلة (قبل الطرح من الهدايا)",
                  "قيمة المطالبة (الهدايا − ناتج المعادلة)", "مؤهلة (عرض مفرق صالح)؟",
                  "عدد الأسطر ضمن الحزمة", "ملاحظة"]

RAW_HEADERS = ["الفاتورة", "التاريخ", "اسم الزبون", "اسم المادة", "عرض المفرق",
               "عرض مميز1", "عرض مميز", "كمية", "الهدايا", "ملاحظة"]

# شيت كل مادة: تفصيل سطر بسطر (الفاتورة/الزبون مرجعي فقط، لا يدخلان أي
# حساب — التجميع الآن على مستوى مادة×عرض مفرق فقط، تصحيح 2026-08-31)
ROW_DETAIL_HEADERS = ["الفاتورة", "التاريخ", "اسم الزبون (مرجعي فقط)", "عرض المفرق",
                       "كمية", "الهدايا", "ناتج المعادلة (الإفرادي)",
                       "الصافي = الهدايا − ناتج المعادلة (السعر الإجمالي)",
                       "مُدرَج بالمجموع النهائي؟", "ملاحظة"]


def _group_row_values(g):
    return [
        g.item, g.retail_offer, float(g.qty), float(g.gifts),
        float(g.formula_result) if g.formula_result is not None else "",
        float(g.claim_value) if g.claim_value is not None else "",
        "نعم" if g.eligible else "لا",
        len(g.row_calcs), g.note,
    ]


def _raw_row_values(r):
    return [
        r.invoice, r.date, r.customer, r.item, r.retail_offer, r.special_offer1, r.special_offer,
        float(r.qty), float(r.gifts), r.note,
    ]


def _row_calc_values(rc):
    r = rc.row
    return [
        r.invoice, r.date, r.customer, r.retail_offer, float(r.qty), float(r.gifts),
        float(rc.formula_result), float(rc.net), "نعم" if rc.included else "لا", r.note,
    ]


def _header(ws, headers):
    ws.append(headers)
    r = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = f"A{r + 1}"


def build_workbook(result, claims_table):
    wb = Workbook()

    ws0 = wb.active
    ws0.title = "الملخص"
    ws0.sheet_view.rightToLeft = True
    ws0.append([
        "المنهجية (إعادة بناء ثانية بطلب المستخدم — آخر تحديث 2026-08-31): تعبئة رقم الفاتورة "
        "الصفري من السطر السابق، ثم فلترة العروض المميزة فقط، ثم تجميع حسب (المادة × عرض المفرق) "
        "فقط (بلا اسم الصيدلية/الزبون إطلاقاً)، ثم داخل كل حزمة: إلغاء أزواج \"م. مبيع\"/المرتجعات "
        "المتطابقة (بنفس الكمية والهدايا، بلا اشتراط الزبون)، ثم دمج كل سطر كميته = صفر مع سطر آخر "
        "لنفس الزبون (بإضافة هداياه إليه) وحذف سطر الكمية=صفر، ثم لكل سطر بمفرده: ناتج المعادلة "
        "= الكمية × (الصغير ÷ الكبير) من \"عرض المفرق\"، والصافي = الهدايا − ناتج المعادلة. "
        "الأسطر ذات الصافي السالب أو الصفري تُحذف بالكامل من الحساب (بطلب المستخدم الصريح، رغم "
        "تعارض جزئي موثَّق مع دليل مرجعي حي — انظر ملاحظة التحقق في compensation/engine.py). "
        "المطالبة النهائية للحزمة = مجموع هدايا الأسطر المُدرَجة − مجموع ناتج معادلتها."
    ])
    ws0.append([])
    _header(ws0, ["إجمالي سطور الحركة الخام", "أسطر بعد فلتر العروض المميزة",
                  "مواد مستبعدة كلياً (بلا عرض مميز)", "فواتير م. مبيع مستبعدة",
                  "أزواج م. مبيع/مبيع محذوفة معاً", "أزواج مرتجع/مبيع محذوفة",
                  "مرتجعات بدون مبيع مطابق", "أسطر كمية=صفر مدموجة بسطر آخر",
                  "أسطر كمية=صفر بلا زبون آخر لدمجها", "أسطر صافٍ سالب/صفري مستبعدة", "عدد المواد",
                  "عدد الحزم (مادة×عرض مفرق)", "حزم بلا عرض مفرق صالح",
                  "إجمالي الكمية (المُدرَجة)", "إجمالي الهدايا (المُدرَجة)",
                  "إجمالي ناتج المعادلة (المُدرَج)", "إجمالي قيمة المطالبة"])
    ws0.append([
        result["raw_row_count"], result["offer_row_count"], result["ignored_items_count"],
        result["excluded_mabee_rows_count"], result["cancelled_mabee_pairs"],
        result["cancelled_pairs"], result["unmatched_returns_count"],
        result["merged_zero_qty_rows"], result["unmerged_zero_qty_rows_count"],
        result["excluded_nonpositive_rows_count"], result["items_count"],
        result["groups_count"], len(result["ineligible_groups"]),
        float(result["total_qty"]), float(result["total_gifts"]),
        float(result["total_formula_result"]), float(result["total_claim_value"]),
    ])
    for i, w in enumerate([20, 20, 20, 18, 20, 20, 18, 20, 20, 20, 14, 20, 18, 18, 18, 22, 20], start=1):
        ws0.column_dimensions[get_column_letter(i)].width = w

    if result["ignored_items"]:
        ws_ig = wb.create_sheet("مواد مستبعدة (بلا عرض مميز)")
        ws_ig.sheet_view.rightToLeft = True
        ws_ig.append([
            "هذه المواد لا تملك أي قيمة \"عرض مميز\" أو \"عرض مميز1\" على أي سطر لها في كل الملف، "
            "فاستُبعدت بالكامل من الحساب بطلب المستخدم الصريح — لا تظهر أرقامها في أي شيت آخر."
        ])
        _header(ws_ig, ["اسم المادة"])
        for item in result["ignored_items"]:
            ws_ig.append([item])
        ws_ig.column_dimensions["A"].width = 46

    ws_items = wb.create_sheet("ملخص حسب المادة")
    ws_items.sheet_view.rightToLeft = True
    _header(ws_items, ["اسم المادة", "عدد الحزم (عرض مفرق)", "مجموع الكمية",
                        "مجموع الهدايا", "ناتج المعادلة (قبل الطرح)",
                        "إجمالي قيمة المطالبة (بعد الطرح)", "المطالبة اليدوية (مرجع)"])
    for item, groups in sorted(result["per_item"].items()):
        qty_sum = sum((g.qty for g in groups), 0)
        gifts_sum = sum((g.gifts for g in groups), 0)
        formula_sum = sum((g.formula_result or 0) for g in groups)
        value_sum = sum((g.claim_value or 0) for g in groups)
        ws_items.append([
            item, len(groups), float(qty_sum), float(gifts_sum), float(formula_sum), float(value_sum),
            float(claims_table.get(item, "")) if claims_table.get(item) is not None else "",
        ])
    for i, w in enumerate([46, 20, 16, 16, 20, 22, 20], start=1):
        ws_items.column_dimensions[get_column_letter(i)].width = w

    # شيت لكل مادة: تفصيل كل سطر بمفرده لكل حزمة (عرض مفرق) — يطابق شكل
    # ملف المستخدم المرجعي (عمودا "الإفرادي" و"السعر الإجمالي")، مع صف
    # إجمالي أصفر في نهاية كل حزمة (وصف إجمالي عام للمادة في النهاية إن
    # تعددت حزم العروض).
    used_names = set()
    for item, groups in sorted(result["per_item"].items()):
        base = safe_sheet_title(item)[:28] or "مادة"
        name = base
        n = 1
        while name in used_names:
            n += 1
            name = f"{base[:25]}_{n}"
        used_names.add(name)
        ws = wb.create_sheet(name)
        ws.sheet_view.rightToLeft = True
        _header(ws, ROW_DETAIL_HEADERS)

        material_qty = material_gifts = material_formula = material_claim = 0
        for g in sorted(groups, key=lambda g: g.retail_offer):
            if not g.eligible:
                ws.append([f"⚠ حزمة بعرض مفرق '{g.retail_offer}' — {g.note}"] + [""] * (len(ROW_DETAIL_HEADERS) - 1))
                for c in range(1, len(ROW_DETAIL_HEADERS) + 1):
                    ws.cell(row=ws.max_row, column=c).fill = FLAG_FILL
                continue
            for rc in g.row_calcs:
                ws.append(_row_calc_values(rc))
                if not rc.included:
                    for c in range(1, len(ROW_DETAIL_HEADERS) + 1):
                        ws.cell(row=ws.max_row, column=c).fill = EXCLUDED_ROW_FILL
            total_row = [f"إجمالي عرض المفرق '{g.retail_offer}'", "", "", "",
                         float(g.qty), float(g.gifts), float(g.formula_result), float(g.claim_value), "", ""]
            ws.append(total_row)
            for c in range(1, len(ROW_DETAIL_HEADERS) + 1):
                cell = ws.cell(row=ws.max_row, column=c)
                cell.fill = TOTAL_FILL
                cell.font = Font(bold=True)
            material_qty += g.qty
            material_gifts += g.gifts
            material_formula += (g.formula_result or 0)
            material_claim += (g.claim_value or 0)

        eligible_count = sum(1 for g in groups if g.eligible)
        if eligible_count > 1:
            ws.append(["إجمالي المادة (كل حزم عروض المفرق)", "", "", "",
                       float(material_qty), float(material_gifts), float(material_formula),
                       float(material_claim), "", ""])
            for c in range(1, len(ROW_DETAIL_HEADERS) + 1):
                cell = ws.cell(row=ws.max_row, column=c)
                cell.fill = TOTAL_FILL
                cell.font = Font(bold=True, size=12)
        for i, w in enumerate([16, 12, 34, 12, 10, 10, 16, 20, 14, 46], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    if result["ineligible_groups"]:
        ws_neg = wb.create_sheet("حزم بلا عرض مفرق صالح")
        ws_neg.sheet_view.rightToLeft = True
        ws_neg.append([
            "هذه حزم (مادة × عرض مفرق) لا تملك قيمة \"عرض مفرق\" صالحة (فارغة/\"-\"/\"0\" أو بصيغة "
            "غير مفهومة)، فتعذّر تطبيق المعادلة عليها — الكمية والهدايا ظاهرة كاملة (خام، بلا حذف "
            "أسطر سالبة/صفرية لعدم وجود معادلة أصلاً)، لكن بلا قيمة مطالبة. راجعها يدوياً."
        ])
        _header(ws_neg, GROUP_HEADERS)
        for g in result["ineligible_groups"]:
            ws_neg.append(_group_row_values(g))
            for c in range(1, len(GROUP_HEADERS) + 1):
                ws_neg.cell(row=ws_neg.max_row, column=c).fill = FLAG_FILL
        for i, w in enumerate([40, 12, 14, 14, 18, 20, 16, 16, 46], start=1):
            ws_neg.column_dimensions[get_column_letter(i)].width = w

    if result["excluded_mabee_rows"]:
        ws_mb = wb.create_sheet("فواتير م. مبيع مستبعدة")
        ws_mb.sheet_view.rightToLeft = True
        ws_mb.append([
            "هذه أسطر فواتير \"م. مبيع\" (تصحيحات/سحوبات مبيعات) بلا سطر بيع مطابق (بنفس الكمية "
            "والهدايا ضمن نفس المادة وعرض المفرق) لإلغائها معه — مستبعدة دائماً من الحساب."
        ])
        _header(ws_mb, RAW_HEADERS)
        for r in result["excluded_mabee_rows"]:
            ws_mb.append(_raw_row_values(r))
            for c in range(1, len(RAW_HEADERS) + 1):
                ws_mb.cell(row=ws_mb.max_row, column=c).fill = FLAG_FILL
        for i, w in enumerate([16, 12, 34, 40, 12, 12, 12, 8, 8, 46], start=1):
            ws_mb.column_dimensions[get_column_letter(i)].width = w

    if result["excluded_nonpositive_rows"]:
        ws_np = wb.create_sheet("أسطر صافٍ سالب أو صفري مستبعدة")
        ws_np.sheet_view.rightToLeft = True
        ws_np.append([
            "هذه أسطر كان صافيها (الهدايا − ناتج المعادلة لهذا السطر وحده) سالباً أو صفراً — حُذفت "
            "بالكامل من الحساب (كميتها وهداياها وناتج معادلتها) بطلب المستخدم الصريح 2026-08-31. "
            "ملاحظة: هذه القاعدة عارضت جزئياً دليلاً مرجعياً حياً لمادة واحدة (اوستيو فيكس) — انظر "
            "التوثيق الكامل في compensation/engine.py قبل اعتماد الأرقام نهائياً."
        ])
        _header(ws_np, ["الفاتورة", "التاريخ", "اسم الزبون", "اسم المادة", "عرض المفرق",
                        "كمية", "الهدايا", "ناتج المعادلة", "الصافي (مستبعد)"])
        for rc in result["excluded_nonpositive_rows"]:
            r = rc.row
            ws_np.append([r.invoice, r.date, r.customer, r.item, r.retail_offer,
                          float(r.qty), float(r.gifts), float(rc.formula_result), float(rc.net)])
            for c in range(1, 10):
                ws_np.cell(row=ws_np.max_row, column=c).fill = EXCLUDED_ROW_FILL
        for i, w in enumerate([16, 12, 34, 40, 12, 10, 10, 14, 14], start=1):
            ws_np.column_dimensions[get_column_letter(i)].width = w

    if result["unmerged_zero_qty_rows"]:
        ws_uz = wb.create_sheet("أسطر كمية=صفر بلا زبون آخر")
        ws_uz.sheet_view.rightToLeft = True
        ws_uz.append([
            "هذه أسطر كميتها = صفر، ولم يوجد لها سطر آخر بنفس الزبون ضمن نفس المادة وعرض المفرق "
            "لدمج هداياها معه — بقيت ظاهرة بمفردها ولم تُحذف صامتة. راجعها يدوياً (ستدخل الحساب "
            "بناتج معادلة = صفر، فيُحتسَب صافيها = هداياها كاملة إن كانت موجَبة)."
        ])
        _header(ws_uz, RAW_HEADERS)
        for r in result["unmerged_zero_qty_rows"]:
            ws_uz.append(_raw_row_values(r))
            for c in range(1, len(RAW_HEADERS) + 1):
                ws_uz.cell(row=ws_uz.max_row, column=c).fill = FLAG_FILL
        for i, w in enumerate([16, 12, 34, 40, 12, 12, 12, 8, 8, 46], start=1):
            ws_uz.column_dimensions[get_column_letter(i)].width = w

    if result["unmatched_returns"]:
        ws_r = wb.create_sheet("مرتجعات بدون مبيع مطابق")
        ws_r.sheet_view.rightToLeft = True
        ws_r.append([
            "هذه أسطر مرتجعات لم يوجد لها سطر بيع مطابق (نفس الكمية والهدايا ضمن نفس المادة وعرض "
            "المفرق) لإلغائها معه — بقيت كما هي ودخلت الحساب كأي سطر آخر (قد تُحذف لاحقاً إن كان "
            "صافيها سالباً أو صفراً كأي سطر عادي). راجعها يدوياً."
        ])
        _header(ws_r, RAW_HEADERS)
        for r in result["unmatched_returns"]:
            ws_r.append(_raw_row_values(r))
            for c in range(1, len(RAW_HEADERS) + 1):
                ws_r.cell(row=ws_r.max_row, column=c).fill = FLAG_FILL
        for i, w in enumerate([16, 12, 34, 40, 12, 12, 12, 8, 8, 46], start=1):
            ws_r.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
