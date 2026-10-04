import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
OK_FILL = PatternFill("solid", fgColor="E7F7EF")
BAD_FILL = PatternFill("solid", fgColor="FDEAEA")
# تصحيح 2026-09-29 (ميزة "تفصيل حسب الشخص" الجديدة، طلب المستخدم الصريح):
#  - DEFICIT_FILL: نفس أحمر BAD_FILL — يمثّل "ناقص" (فرق سالب: الدائن أقل
#    من المدين، أي أن على الموزّع مبلغاً لم يُسدَّد كشفه بالكامل).
#  - EXCESS_FILL: أزرق فاتح مميَّز — "زائد" (فرق موجب: دُفع أكثر من قيمة
#    الكشف)، حتى يميّزه القارئ فوراً عن حالة "ناقص" رغم أن كليهما "غير مطابق".
#  - SHARED_FILL: كهرماني، نفس عائلة FLAG_FILL المستخدمة في وحدة عمولات
#    الموزعين لعلامات "تحتاج انتباه" — يمثّل هنا كشفاً "مشترك" بين أكثر من
#    موزّع (لا يُحسَب ضمن مجموع أي منهما الفردي، لكن يجب أن يظهر بوضوح عندهم).
DEFICIT_FILL = BAD_FILL
EXCESS_FILL = PatternFill("solid", fgColor="DBEAFE")
SHARED_FILL = PatternFill("solid", fgColor="FDE68A")
UNLINKED_FILL = PatternFill("solid", fgColor="E5E7EB")


def _header(ws, headers):
    ws.append(headers)
    r = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = f"A{r + 1}"


def _identifier_label(g):
    if g["statement_numbers"]:
        return ", ".join(g["statement_numbers"])
    if g["statement_names"]:
        return ", ".join(g["statement_names"])
    return "—"


def _identifier_kind_label(g):
    if g["statement_numbers"]:
        return "رقم"
    if g["statement_names"]:
        return "اسم (مطابقة تقريبية)"
    return "—"


def build_workbook(result, meta, person_breakdown=None):
    wb = Workbook()
    person_breakdown = person_breakdown or {}

    ws0 = wb.active
    ws0.title = "الملخص"
    ws0.sheet_view.rightToLeft = True
    s = result["summary"]
    for row in [
        ["الملف", meta.get("file_name", "")],
        ["إجمالي عدد الحركات", s["total_rows"]],
        ["حركات مرتبطة بمعرّف كشف (رقم أو اسم)", s["linked_rows"]],
        ["حركات بلا معرّف كشف (تحتاج مراجعة يدوية)", s["unlinked_rows"]],
        ["عدد مجموعات الكشوفات", s["groups_count"]],
        ["  منها: مرتبطة بالاسم فقط (بلا رقم كشف)", s["name_linked_groups"]],
        ["مجموعات متطابقة", s["matched_groups"]],
        ["مجموعات غير متطابقة", s["mismatched_groups"]],
        ["  منها: فردية (موزّع واحد)", s["solo_mismatched_groups"]],
        ["  منها: مشتركة بين أكثر من موزّع (مشترك)", s["shared_mismatched_groups"]],
        ["إجمالي فرق المجموعات غير المتطابقة", float(s["mismatched_amount"])],
        ["عدد الأشخاص (الموزّعين) الظاهرين في الملف", len(person_breakdown)],
    ]:
        ws0.append(row)
    ws0.column_dimensions["A"].width = 46
    ws0.column_dimensions["B"].width = 24

    ws1 = wb.create_sheet("المطابقة حسب الكشف")
    ws1.sheet_view.rightToLeft = True
    _header(ws1, [
        "معرّف الكشف", "نوع المعرّف", "مشترك بين أكثر من موزّع؟", "الأشخاص (الموزّعون)",
        "الحسابات الفرعية", "إجمالي مدين", "إجمالي دائن (مجمّع)", "الفرق", "الحالة", "عدد الحركات",
    ])
    for g in result["groups"]:
        ws1.append([
            _identifier_label(g),
            _identifier_kind_label(g),
            "نعم — مشترك" if g["is_shared"] else "لا",
            "، ".join(g["persons"]) or "—",
            ", ".join(g["subaccounts"]),
            float(g["total_debit"]), float(g["total_credit"]), float(g["difference"]),
            "مطابق" if g["matched"] else "غير مطابق", len(g["rows"]),
        ])
        r = ws1.max_row
        if g["is_shared"]:
            ws1.cell(row=r, column=3).fill = SHARED_FILL
        ws1.cell(row=r, column=9).fill = OK_FILL if g["matched"] else BAD_FILL
    for i, w in enumerate([20, 20, 20, 30, 30, 16, 18, 14, 14, 12], start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("تفاصيل الحركات")
    ws2.sheet_view.rightToLeft = True
    _header(ws2, ["معرّف الكشف", "الشخص (بيد)", "التاريخ", "رقم السند", "البيان", "الحساب الفرعي", "مدين", "دائن"])
    for g in result["groups"]:
        ident = _identifier_label(g)
        for r in g["rows"]:
            ws2.append([ident, r["person"] or "—", r["date"], r["voucher_no"], r["narration"],
                        r["subaccount"], float(r["debit"]), float(r["credit"])])
    for i, w in enumerate([18, 20, 12, 12, 55, 26, 14, 14], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ---- تفصيل حسب الشخص (بيد) — طلب المستخدم الصريح 2026-09-29 ----
    ws3 = wb.create_sheet("ملخص حسب الشخص (بيد)")
    ws3.sheet_view.rightToLeft = True
    _header(ws3, [
        "الشخص (بيد)", "إجمالي الناقص", "إجمالي الزائد", "الصافي",
        "عدد الكشوفات الناقصة", "عدد الكشوفات الزائدة",
        "عدد الكشوفات المشتركة (غير محتسبة بالصافي)", "حركات بلا معرّف كشف",
    ])
    for name, b in person_breakdown.items():
        ws3.append([
            name, float(b["deficit_total"]), float(b["excess_total"]), float(b["net_total"]),
            len(b["deficit_groups"]), len(b["excess_groups"]), len(b["shared_groups"]), len(b["unlinked_rows"]),
        ])
        r = ws3.max_row
        if b["deficit_total"] != 0:
            ws3.cell(row=r, column=2).fill = DEFICIT_FILL
        if b["excess_total"] != 0:
            ws3.cell(row=r, column=3).fill = EXCESS_FILL
    for i, w in enumerate([26, 16, 16, 16, 16, 16, 22, 18], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    ws4 = wb.create_sheet("تفصيل حسب الشخص (بيد)")
    ws4.sheet_view.rightToLeft = True
    _header(ws4, [
        "الشخص (بيد)", "النوع", "معرّف الكشف", "نوع المعرّف", "الحسابات الفرعية",
        "إجمالي مدين", "إجمالي دائن", "الفرق", "عدد الحركات",
    ])
    for name, b in person_breakdown.items():
        for g in b["deficit_groups"]:
            ws4.append([name, "ناقص", _identifier_label(g), _identifier_kind_label(g),
                        ", ".join(g["subaccounts"]), float(g["total_debit"]), float(g["total_credit"]),
                        float(g["difference"]), len(g["rows"])])
            ws4.cell(row=ws4.max_row, column=2).fill = DEFICIT_FILL
        for g in b["excess_groups"]:
            ws4.append([name, "زائد", _identifier_label(g), _identifier_kind_label(g),
                        ", ".join(g["subaccounts"]), float(g["total_debit"]), float(g["total_credit"]),
                        float(g["difference"]), len(g["rows"])])
            ws4.cell(row=ws4.max_row, column=2).fill = EXCESS_FILL
        for g in b["shared_groups"]:
            ws4.append([name, "مشترك (غير محتسب بالصافي)", _identifier_label(g), _identifier_kind_label(g),
                        ", ".join(g["subaccounts"]), float(g["total_debit"]), float(g["total_credit"]),
                        float(g["difference"]), len(g["rows"])])
            ws4.cell(row=ws4.max_row, column=2).fill = SHARED_FILL
    for i, w in enumerate([26, 26, 20, 20, 30, 16, 16, 14, 12], start=1):
        ws4.column_dimensions[get_column_letter(i)].width = w

    ws5 = wb.create_sheet("حركات بلا معرّف كشف")
    ws5.sheet_view.rightToLeft = True
    _header(ws5, ["الشخص (بيد)", "التاريخ", "رقم السند", "البيان", "الحساب الفرعي", "مدين", "دائن"])
    for r in result["unlinked"]:
        ws5.append([r["person"] or "—", r["date"], r["voucher_no"], r["narration"],
                    r["subaccount"], float(r["debit"]), float(r["credit"])])
        ws5.cell(row=ws5.max_row, column=1).fill = UNLINKED_FILL
    for i, w in enumerate([20, 12, 12, 55, 26, 14, 14], start=1):
        ws5.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
