import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _header(ws, headers):
    ws.append(headers)
    r = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = f"A{r + 1}"


def build_workbook(combined_rows, discount_rows, summary, meta):
    wb = Workbook()

    ws0 = wb.active
    ws0.title = "الملخص"
    ws0.sheet_view.rightToLeft = True
    ws0.append(["إجمالي عدد السطور", summary["total_rows"]])
    for src, count in summary["by_source"].items():
        ws0.append([f"   عدد سطور ({src})", count])
    ws0.append(["إجمالي القيمة المؤجلة", float(summary["total_deferred_value"])])
    ws0.append(["عدد سطور الحسم الممنوح", summary["discount_rows"]])
    ws0.append(["إجمالي الحسم الممنوح (سالب)", float(summary["discount_total"])])
    ws0.column_dimensions["A"].width = 34
    ws0.column_dimensions["B"].width = 20

    ws1 = wb.create_sheet("تجميع (مبيعات+مرتجعات+ذمم)")
    ws1.sheet_view.rightToLeft = True
    _header(ws1, ["الفاتورة", "اسم الزبون", "مركز الكلفة", "القيمة المؤجلة", "تاريخ التسليم", "الكتلة"])
    for r in combined_rows:
        ws1.append([r["invoice"], r["customer"], r["cost_center"], float(r["deferred_value"]),
                    r["delivery_date"], r["zone"]])
    for i, w in enumerate([12, 46, 26, 16, 14, 20], start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("الحسم الممنوح (دفعات)")
    ws2.sheet_view.rightToLeft = True
    _header(ws2, ["رقم السند", "مدين (سالب)", "التاريخ", "مركز الكلفة", "الحساب المقابل", "البيان"])
    for r in discount_rows:
        ws2.append([r["voucher_no"], float(r["debit"]), r["date"], r["cost_center"], r["counterpart"], r["narration"]])
    for i, w in enumerate([12, 16, 12, 26, 30, 55], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
