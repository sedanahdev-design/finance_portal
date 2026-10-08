"""محرك عمولات الموزعين الخارجيين (الفكرة التاسعة).

لكل موزّع خارجي ملف "سحب من العميل" (دفتر أستاذ) يحوي:
    - مدين  = كشوفات (Statements) بذمّة الموزّع
    - دائن  = دفعات و/أو مرتجعات محصّلة لصالح الموزّع (مجمّعة معاً في
              عمود دائن واحد داخل دفتر الأستاذ الخام، دون تمييز مباشر)

--- تعديل 2026-08-26 (الأول) --------------------------------------------
المعادلة القديمة: العمولة = ((مبيع + مرتجع) − كشوفات) × 3%
المعادلة بعده:   العمولة = (مرتجع − كشوفات) × 3%  — أُسقط بند "مبيع".

--- تعديل 2026-10-08 (الثاني — بطلب المستخدم الصريح، فيشة "شادي حسن") ---
المستخدم رفع فيشة حساب يدوية حقيقية لموزّع "شادي حسن" (شهر 9) بصيغة
مختلفة قليلاً عن الفرضية السابقة: لا يوجد فيها بند "مبيع" إطلاقاً، بل
"الدفعات" و"مرتجعات" فقط، والمعادلة فيها ظاهرة نصاً:
    العمولة = (الكشوفات − مرتجعات) × 3%
أي أن اتجاه الطرح انعكس (كشوفات ناقص مرتجعات، لا مرتجعات ناقص كشوفات)،
وبند "مبيع (مستبعد من العمولة)" أصبح "دفعة" (Payment) — لأن حركات الدائن
غير المرتجع في هذا السياق هي دفعات مالية من الموزّع، لا مبيعات. أكّد
المستخدم عبر أسئلة توضيحية:
  • المعادلة: "(كشوفات − مرتجعات) × 3% — متل الفيشة الجديدة"
  • اتجاه الرصيد: "الكشوفات تزيد الرصيد (رصيد = سابق + كشوفات − دفعات
    − مرتجعات)"
  • مصدر البيانات: نفس ملف دفتر الأستاذ الحالي (مدين/دائن) — لا ملف جديد.
  • النطاق: الفيشة الجديدة تُضاف كشكل إضافي (شيت مستقل) جنب شيت
    المدين/دائن الموجود، ولا تستبدله.

المعادلة الحالية (بعد هذا التعديل):
    العمولة = (كشوفات − مرتجعات) × 3%
    أي: commission_base = total_debit − total_returns

اتجاه الرصيد الحالي (بعد هذا التعديل):
    الرصيد الختامي = الرصيد الافتتاحي + كشوفات − (دفعات + مرتجعات)
    أي: closing_balance = opening_balance + total_debit − total_credit
    (total_credit = total_payments + total_returns، فتُطرح معاً كما في
    معادلة المستخدم الحرفية "سابق + كشوفات − دفعات − مرتجعات").
    وللاتساق، انعكس أيضاً اتجاه حساب الرصيد الافتتاحي من صف "الرصيد
    السابق" نفسه: opening_balance = مدين_السطر − دائن_السطر (بدل العكس)،
    حتى يبقى المدين موجِباً للرصيد والدائن سالباً له في كل مكان بالملف.

مصدر "مرتجع" (لم يتغيّر): بما أن دفتر الأستاذ الخام لا يميّز داخلياً بين
"دائن-دفعة" و"دائن-مرتجع" (كلاهما في عمود دائن واحد)، تُستخرج المرتجعات
من نص عمود "البيان" الحر: أي حركة دائن يحتوي بيانها على كلمة "مرتجع"
تُحتسب ضمن المرتجعات (أساس العمولة)؛ وأي حركة دائن أخرى (بيان فارغ أو
لا يحوي "مرتجع") تُعامَل كـ"دفعة" (كان اسمها "مبيع" سابقاً) وتُستبعد من
أساس العمولة — لكنها لا تُحذف من البيانات: تبقى ظاهرة في صفوف الحركات
(كل حركة دائن موسومة بنوعها: مرتجع / دفعة) وضمن إجمالي منفصل
"total_payments" (بدّل اسم "total_sales_excluded" السابق) في نتيجة كل
موزّع، حتى لا يُخفى أي رقم عن المستخدم.

الرصيد الافتتاحي ("الرصيد السابق") رصيد دوّار مرحّل من الشهر السابق، ولا
يدخل في حساب عمولة الشهر الحالي، لكنه يظهر في التقرير لإظهار الرصيد
الختامي الصحيح.

فيشة "كشف حساب" الجديدة (شيت إضافي مستقل لكل موزّع في excel_export.py):
لإخراج "رقم الايصال" (للدفعات) و"اسم الصيدلية" (للمرتجعات) — وهما غير
موجودين كعمودين مستقلّين في دفتر الأستاذ الخام — تُستخرَجان اجتهاداً من
نص عمود "البيان" الحر: رقم يتبع كلمة "إيصال" للدفعات، والنص بعد كلمة
"مرتجع" (بعد حذف الفواصل) للمرتجعات. كذلك تُستخرج فترة الكشف (الشهر،
وتاريخ الرصيد الافتتاحي/الختامي) من سطر "اعتباراً من ... ولغاية ..." في
رأس الملف الخام إن وُجد. هذه اجتهادات معقولة لم تُختبر بعد على ملف دفتر
أستاذ حقيقي لموزّع خارجي (لم يتوفر وقت هذا التعديل) — يلزم ملف حقيقي
واحد على الأقل للتحقق النهائي من مطابقتها للفيشة اليدوية بالضبط.
"""

import re
from datetime import date as _date, timedelta
from decimal import Decimal

import openpyxl

COMMISSION_RATE = Decimal("0.03")

HEADER_SIGNATURE = {"رقم السند", "مدين", "دائن", "التاريخ"}
STOP_MARKERS = {"المجموع", "الرصيد النهائي"}

# كلمة تُستخدم لاستخراج حركات "مرتجع" من نص عمود "البيان" الحر ضمن حركات
# الدائن. أي حركة دائن لا يحتوي بيانها على هذه الكلمة تُعتبر "دفعة"
# وتُستبعد من أساس العمولة.
RETURN_MARKER = "مرتجع"

# لاستخراج "رقم الايصال" (لحركات الدفعة) و"اسم الصيدلية" (لحركات
# المرتجع) من نص "البيان" الحر — راجع توثيق تعديل 2026-10-08 أعلاه.
_RECEIPT_NO_RE = re.compile(r"إيصال\s*(\d+)")
_PERIOD_RE = re.compile(r"اعتبار\S*\s*من\s*([\d]{1,2}-[\d]{1,2}-[\d]{4})\s*ولغاية\s*([\d]{1,2}-[\d]{1,2}-[\d]{4})")


def _clean(v):
    if v is None:
        return ""
    return str(v).strip() if isinstance(v, str) else v


def _to_decimal(v):
    if v is None or v == "":
        return Decimal("0")
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v)).quantize(Decimal("0.01"))
    s = str(v).replace(",", "").strip()
    m = re.match(r"^-?\d+(\.\d+)?", s)
    if not m:
        return Decimal("0")
    try:
        return Decimal(m.group(0)).quantize(Decimal("0.01"))
    except Exception:  # noqa: BLE001
        return Decimal("0")


def _distributor_name_from_account(all_rows):
    for row in all_rows[:6]:
        if row and row[0] and "اسم الحساب" in str(row[0]):
            m = re.search(r"-(.+)$", str(row[0]))
            if m:
                return m.group(1).strip()
    return None


def _parse_dm_date(s):
    """يحوّل نصاً مثل '23-9-2026' (يوم-شهر-سنة) إلى date، أو None."""
    parts = str(s).strip().split("-")
    if len(parts) != 3:
        return None
    try:
        d, m, y = int(parts[0]), int(parts[1]), int(parts[2])
        return _date(y, m, d)
    except Exception:  # noqa: BLE001
        return None


def _parse_period(all_rows):
    """يستخرج فترة الكشف (بداية/نهاية) من سطر 'اعتباراً من ... ولغاية ...'
    في رأس الملف الخام، إن وُجد. يُستخدم فقط لتزيين فيشة 'كشف الحساب'
    (تسمية الشهر وتاريخي الرصيد الافتتاحي/الختامي) — لا يدخل في أي حساب
    مالي. يعيد (period_start, period_end) كـ date أو (None, None)."""
    for row in all_rows[:8]:
        if not row:
            continue
        for cell in row:
            if not cell or not isinstance(cell, str):
                continue
            m = _PERIOD_RE.search(cell)
            if m:
                return _parse_dm_date(m.group(1)), _parse_dm_date(m.group(2))
    return None, None


def _extract_receipt_no(reference):
    """رقم الايصال لحركة 'دفعة' — يُستخرج من نص البيان إن وُجدت كلمة
    'إيصال' متبوعة برقم (راجع توثيق تعديل 2026-10-08). قد يعود فارغاً
    (وهو الحال الغالب فعلياً في الفيشة المرجعية)."""
    m = _RECEIPT_NO_RE.search(str(reference))
    return m.group(1) if m else ""


def _extract_pharmacy_name(reference):
    """اسم الصيدلية لحركة 'مرتجع' — يُستخرج من نص البيان بعد حذف كلمة
    'مرتجع' والفواصل المحيطة بها (راجع توثيق تعديل 2026-10-08). إن لم
    يبقَ نص بعد الحذف، يُعاد البيان الكامل كما هو (أفضل من تفريغه بصمت)."""
    s = str(reference)
    idx = s.find(RETURN_MARKER)
    if idx == -1:
        return s.strip()
    rest = s[idx + len(RETURN_MARKER):]
    rest = rest.strip(" -:_،")
    return rest or s.strip()


def parse_distributor_ledger(file_obj, display_name=None):
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    name = display_name or _distributor_name_from_account(all_rows) or "موزّع غير محدد"
    period_start, period_end = _parse_period(all_rows)

    header_idx = header = None
    for i, row in enumerate(all_rows):
        cells = {str(_clean(c)) for c in row if c is not None}
        if HEADER_SIGNATURE.issubset(cells):
            header_idx, header = i, [str(_clean(c)) for c in row]
            break
    if header_idx is None:
        raise ValueError(f"لم يتم العثور على صف العناوين المتوقع داخل ملف {name}.")

    col = {h: header.index(h) for h in header if h}

    def get(row, key):
        idx = col.get(key)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    opening_balance = Decimal("0")
    debit_rows, credit_rows = [], []
    closing_balance_from_source = None

    for row in all_rows[header_idx + 1:]:
        if row is None or all(c in (None, "") for c in row):
            continue
        first_cell = _clean(row[0])
        if str(first_cell) in STOP_MARKERS:
            # صف الإجمالي/الرصيد النهائي له بنية مختلفة عن باقي الصفوف
            # (بلا رقم سند): [التسمية، مجموع مدين، مجموع دائن، الرصيد، ...]
            if len(row) > 3:
                closing_balance_from_source = _to_decimal(row[3])
            if first_cell == "الرصيد النهائي":
                break
            continue
        # تصحيح 2026-10-08 (ضمن هذا التعديل): صف "الرصيد السابق" يحمل
        # هذا النص داخل عمود "البيان"، لا العمود الأول (رقم السند، الذي
        # يكون صفراً فيه عادة) — الفحص القديم كان يقارن العمود الأول
        # فيفوت هذا الصف دوماً على بيانات حقيقية، فيبقى opening_balance
        # صفراً بصمت. تأكّدنا من ترتيب الأعمدة هذا على ملف دفتر أستاذ
        # حقيقي بنفس بنية الرأس (رقم السند/البيان/مدين/دائن...).
        if _clean(get(row, "البيان")) == "الرصيد السابق":
            debit0 = _to_decimal(get(row, "مدين"))
            credit0 = _to_decimal(get(row, "دائن"))
            # تعديل 2026-10-08: اتجاه الرصيد الآن "كشوفات تزيد" — انظر
            # توثيق أعلى الملف. مدين موجِب للرصيد، دائن سالب له.
            opening_balance = debit0 - credit0
            continue

        debit = _to_decimal(get(row, "مدين"))
        credit = _to_decimal(get(row, "دائن"))
        if debit == 0 and credit == 0:
            continue

        raw_ref = get(row, "البيان")
        reference = _clean(raw_ref)
        voucher_no = _clean(get(row, "رقم السند"))
        date = _clean(get(row, "التاريخ"))

        if debit:
            debit_rows.append({
                "voucher_no": voucher_no, "reference": reference, "date": date,
                "amount": debit,
            })
        if credit:
            # هل هذه الحركة "مرتجع" بحسب نص البيان؟ أي حركة دائن أخرى
            # تُعامَل كـ"دفعة" (كانت تُسمّى "مبيع" قبل تعديل 2026-10-08)
            # وتُستبعد من أساس العمولة، مع بقائها ظاهرة هنا موسومة.
            is_return = RETURN_MARKER in str(reference)
            row_data = {
                "voucher_no": voucher_no, "reference": reference, "date": date,
                "amount": credit,
                "kind": "مرتجع" if is_return else "دفعة",
                "is_return": is_return,
                "receipt_no": "" if is_return else _extract_receipt_no(reference),
                "pharmacy_name": _extract_pharmacy_name(reference) if is_return else "",
            }
            credit_rows.append(row_data)

    total_debit = sum((r["amount"] for r in debit_rows), Decimal("0"))
    total_credit = sum((r["amount"] for r in credit_rows), Decimal("0"))
    total_returns = sum((r["amount"] for r in credit_rows if r["is_return"]), Decimal("0"))
    # "total_payments" بدّل اسم "total_sales_excluded" السابق (تعديل
    # 2026-10-08) — نفس الحساب: إجمالي حركات الدائن غير-المرتجع.
    total_payments = total_credit - total_returns

    # المعادلة الحالية (تعديل 2026-10-08): العمولة = (كشوفات − مرتجعات) × 3%
    commission_base = total_debit - total_returns
    commission = (commission_base * COMMISSION_RATE).quantize(Decimal("0.01"))

    # الرصيد الختامي (تعديل 2026-10-08): سابق + كشوفات − (دفعات + مرتجعات)
    closing_balance = closing_balance_from_source if closing_balance_from_source is not None else (
        opening_balance + total_debit - total_credit
    )

    # تسميات تاريخية لفيشة "كشف الحساب" فقط (لا تدخل في أي حساب) — راجع
    # توثيق 2026-10-08 أعلى الملف بخصوص الاجتهاد في استخراجها.
    month_number = period_end.month if period_end else (
        period_start.month if period_start else None
    )
    opening_date_label = f"{(period_start - timedelta(days=1)).day}-{(period_start - timedelta(days=1)).month}" if period_start else ""
    closing_date_label = f"{period_end.day}-{period_end.month}" if period_end else ""

    return {
        "name": name,
        "opening_balance": opening_balance,
        "debit_rows": debit_rows,
        "credit_rows": credit_rows,
        "total_debit": total_debit,
        "total_credit": total_credit,
        "total_returns": total_returns,
        "total_payments": total_payments,
        "commission_base": commission_base,
        "commission_rate": COMMISSION_RATE,
        "commission": commission,
        "closing_balance": closing_balance,
        "month_number": month_number,
        "opening_date_label": opening_date_label,
        "closing_date_label": closing_date_label,
    }


def build_summary(distributors):
    return {
        "distributors_count": len(distributors),
        "total_commission": sum((d["commission"] for d in distributors), Decimal("0")),
        "total_debit": sum((d["total_debit"] for d in distributors), Decimal("0")),
        "total_credit": sum((d["total_credit"] for d in distributors), Decimal("0")),
        "total_returns": sum((d["total_returns"] for d in distributors), Decimal("0")),
        "total_payments": sum((d["total_payments"] for d in distributors), Decimal("0")),
    }
