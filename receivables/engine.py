"""محرك مطابقة الذمم مع قسم التوزيع (الفكرة الخامسة).

كل حركة في دفتر الأستاذ يُشار في نص "البيان" الخاص بها إلى معرّف "كشف"
مرتبط بها — إما رقماً (أو عدة أرقام دفعة واحدة) أو، حين لا يوجد رقم، اسماً
(اسم صيدلية/مشفى/جهة) وفق طلب المستخدم الصريح 2026-09-29:
"اذا كان عنا رقم الكشف موجود ضمن حقل البيان لازم يكون بعد كلمة كشف رقم
(الرقم) اما اذا ما كان في رقم للكشف ف رح يكون عنا اسم الكشف (اسم الصيدلية)".

أمثلة فعلية من ملف مرجعي حقيقي ("ذمم لدى قسم التوزيع"، تموز 2026):
    "... كشف رقم 3310يومية 1.9.2026"                    -> رقم 3310
    "... كشف 3056-3058-3057  يومية ..."                  -> أرقام 3056,3057,3058
    "... تتمة كشف 3457  يومية ..."                       -> رقم 3457 (يُدمَج مع
                                                              نفس مجموعة 3457،
                                                              "تتمة" لا تعني شيئاً
                                                              مستقلاً)
    "... كشف مشفى المدينة   يومية ..."                   -> اسم "مشفى المدينة"
    "... ص. حلا الصبح يومية ..."                          -> اسم "حلا الصبح"
      (لا توجد كلمة "كشف" إطلاقاً هنا؛ "ص." اختصار "صيدلية" يُستخدم بديلاً)

المطابقة المطلوبة: تجميع كل الحركات (مديناً كانت أم دائناً) التي تشترك في
نفس معرّف الكشف (رقماً كان أم اسماً) ضمن "مجموعة مطابقة" واحدة (خوارزمية
Union-Find، تماماً كالتصميم الأصلي لكن مُوسَّعة لتغطي الأسماء أيضاً)،
ومقارنة إجمالي الدائن بإجمالي المدين على مستوى المجموعة كاملة — بما أن دفعة
تسديد واحدة (أو حركة مدينة واحدة) قد تغطي عدة أرقام/حركات كشف دفعة واحدة،
فلا يمكن فصل مبلغها على كل رقم على حدة بدقة.

اكتشاف مهم أثناء التحقق من ملف مرجعي كامل (طلب المستخدم توضيحه صراحة
2026-09-29): بعض أرقام الكشف تظهر فعلياً تحت أكثر من "شخص" (موزّع) واحد في
نفس الوقت — مثال رقم 3310 الذي ظهر تحت حسابَي "مصطفى سعيد" و"ربيع عبود"
معاً (سند تسديد واحد مقسوم على ذمة موزّعين مختلفين). قرار المستخدم الصريح:
هذه المجموعات "المشتركة" تظهر عند **كل** موزّع معني بوسم "مشترك" واضح، لكن
مبلغ فرقها **لا يُحسَب** ضمن المجموع الفردي الصافي لأي موزّع بعينه (تجنّباً
لتقسيم افتراضي تعسفي لا يعكس واقع البيانات).

"بيد مين هي الذمة": طلب المستخدم أخذ الاسم الأول فقط من "بيد السيد
<اسم1> + <اسم2> ...". تبيّن من الفحص الفعلي أن الاسم الأول المذكور في نص كل
حركة يطابق دائماً (في هذا الملف المرجعي) اسم "الحساب الفرعي" (subaccount)
الذي تقع الحركة تحته في دفتر الأستاذ (كل حساب فرعي = ذمة موزّع واحد بعينه،
عنوانه على الشكل "<رقم>-ذمم [بيد] <اسم الموزّع>") — وهذا مصدر أوثق من محاولة
استخراج الاسم الأول حرفياً من نص كل حركة على حدة، لأن الأسماء المفردة داخل
نص البيان قد تكون مختصرة بشكل غامض (مثال: "بيد السيد مصطفى + سليم..." — لا
يمكن حسم مصطفى سعيد أم مصطفى زيتون من النص وحده، بينما الحساب الفرعي يحسمها
فوراً)، وقد تحوي أخطاء إملائية أو ترتيباً مقلوباً (حالات نادرة رُصدت فعلياً
حيث يظهر الاسم بعد كلمة "إيصال" بدل قبلها، بسبب خلل إدخال). لذلك يُعتمَد اسم
الحساب الفرعي كمصدر "بيد مين" الرسمي — وهو بالضبط "الاسم الأول" الذي طلبه
المستخدم، لأن كل حساب فرعي يمثّل شخصاً واحداً فقط بالتعريف.
"""

import re
import difflib
from decimal import Decimal

import openpyxl

# --- استخراج رقم/اسم الكشف من نص البيان (راجع توثيق أعلى الملف) ---
YAWMIYA_RE = r"(?:يوم[يوة]{0,2}ة|يومية|يويمة)"
KASHF_RE = re.compile(r"كشف\s*(?:رقم)?\s*[:\-]?\s*(.*?)(?=" + YAWMIYA_RE + r"|$)")
SAYD_RE = re.compile(r"ص\s*[.,]\s*(.*?)(?=" + YAWMIYA_RE + r"|$)")
TRAILING_DATE_RE = re.compile(r"\s*\d{1,2}\s*[.\-/]\s*\d{1,2}\s*[.\-/]\s*\d{2,4}\.?\s*$")
NUMBER_SPLIT = re.compile(r"[^\d]+")
PURE_NUMERIC_CHUNK_RE = re.compile(r"^[\d\s\-،,*]+$")

# القديمة، مُبقاة للتوافق الخلفي فقط (لم تعد تُستخدم داخلياً في reconcile؛
# extract_statement_key أدناه أشمل: تدعم أيضاً حالة "لا يوجد رقم" بالاسم).
STATEMENT_PATTERN = re.compile(r"كشف\s*(?:رقم)?\s*[:\-]?\s*([\d][\d\-،,\s]*\d|\d)")


def _strip_tatimma(text):
    """"تتمة كشف" لا تعني شيئاً مستقلاً — إشارة إلى إتمام دفعة كشف (رقم أو
    اسم) مذكور بجانبها مباشرة، وقد ترد قبل المعرّف أو بعده."""
    text = re.sub(r"^\s*تتمة\s*كشف\s*(?:رقم)?\s*[:\-]?\s*", "", text)
    text = re.sub(r"\s*تتمة\s*كشف\s*$", "", text)
    text = re.sub(r"^\s*تتمة\s+", "", text)
    return text.strip()


def _strip_sayd_prefix(text):
    return re.sub(r"^ص\s*[.,]\s*", "", text).strip()


def extract_statement_key(narration):
    """يُعيد (kind, value):
      ("number", {"3056", "3057", ...})  — رقم/أرقام كشف فعلية
      ("name", "مشفى المدينة")            — لا يوجد رقم، فاسم الكشف بديلاً
      (None, None)                        — لا رقم ولا اسم يمكن استخلاصه
    """
    bayan = narration or ""
    text = None
    m = KASHF_RE.search(bayan)
    if m:
        text = m.group(1)
    else:
        m2 = SAYD_RE.search(bayan)
        if m2:
            text = m2.group(1)
    if text is None:
        return (None, None)
    text = TRAILING_DATE_RE.sub("", text).strip()
    text = text.strip(" -:،,")
    text = _strip_tatimma(text)
    text = text.strip(" -:،,")
    if not text:
        return (None, None)
    if PURE_NUMERIC_CHUNK_RE.match(text):
        numbers = set()
        for part in NUMBER_SPLIT.split(text):
            if part and part.isdigit():
                numbers.add(part)
        # تصحيح 2026-09-29 (اكتُشف أثناء التحقق من الملف المرجعي الحقيقي):
        # "0" ليس رقم كشف حقيقياً أبداً — لوحظ فعلياً كقيمة افتراضية متكررة
        # في نص البيان حين لا يُدخَل رقم فعلي (نمط "إيصال 0إيصال مالية
        # 0كشف رقم 0"، 52 حركة موزّعة على موزّعين مختلفين في الملف المرجعي).
        # معاملتها كرقم كشف حقيقي كانت تدمج حركات موزّعين مختلفين تماماً
        # ضمن مجموعة "مشتركة" وهمية لا أساس لها. تُستبعَد هنا فتُعامَل هذه
        # الحركات كحركات "بلا معرّف كشف" (مثل حالة None تماماً).
        numbers.discard("0")
        if numbers:
            return ("number", numbers)
        return (None, None)
    text = _strip_sayd_prefix(text).strip(" -:،,")
    return ("name", text) if text else (None, None)


def extract_statement_numbers(narration):
    """للتوافق الخلفي: أرقام الكشف فقط (بلا حالة الاسم)."""
    kind, val = extract_statement_key(narration)
    return val if kind == "number" else set()


# --- تحديد "الحساب الفرعي" (= الموزّع صاحب الذمة) من عناوين دفتر الأستاذ ---
HEADER_SIGNATURE = {"رقم السند", "مدين", "دائن", "التاريخ"}
STOP_MARKERS = {"المجموع", "الرصيد النهائي"}
SKIP_MARKERS = {"الرصيد السابق"}
SUBACCOUNT_PERSON_RE = re.compile(r"^[\d\-]*-\s*ذمم\s*(?:بيد)?\s*(.+)$")


def _clean(v):
    if v is None:
        return ""
    return str(v).strip() if isinstance(v, str) else v


def _to_decimal(v):
    if v is None or v == "":
        return Decimal("0")
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v)).quantize(Decimal("0.01"))
    s = str(v).replace(",", "").strip()
    m = re.match(r"^-?\d+(\.\d+)?", s)
    if not m:
        return Decimal("0")
    try:
        return Decimal(m.group(0)).quantize(Decimal("0.01"))
    except Exception:  # noqa: BLE001
        return Decimal("0")


def parse_ledger(file_obj):
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header_idx = header = None
    for i, row in enumerate(all_rows):
        cells = {str(_clean(c)) for c in row if c is not None}
        if HEADER_SIGNATURE.issubset(cells):
            header_idx, header = i, [str(_clean(c)) for c in row]
            break
    if header_idx is None:
        raise ValueError("لم يتم العثور على صف العناوين المتوقع داخل الملف.")

    col = {name: header.index(name) for name in header if name}

    def get(row, name):
        idx = col.get(name)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    current_subaccount = ""
    current_person = ""
    rows = []
    for row in all_rows[header_idx + 1:]:
        if row is None or all(c in (None, "") for c in row):
            continue
        first_cell = _clean(row[0])
        narration = _clean(get(row, "البيان"))
        if str(first_cell) in STOP_MARKERS:
            break

        debit = _to_decimal(get(row, "مدين"))
        credit = _to_decimal(get(row, "دائن"))
        raw_date = get(row, "التاريخ")

        # صفوف عناوين/إجماليات الحسابات الفرعية: رقم السند = 0 والتاريخ = 0 (ليست حركة مؤرَّخة فعلية)
        is_header_or_total_row = (
            str(first_cell) in ("0", "")
            and (raw_date in (0, "0", None, "") or "-" not in str(raw_date))
        )
        if is_header_or_total_row:
            if narration and narration not in SKIP_MARKERS and re.match(r"^[\d\-]*-", narration):
                m = SUBACCOUNT_PERSON_RE.match(narration)
                if m:
                    current_subaccount = narration
                    current_person = m.group(1).strip()
                else:
                    # حساب رئيسي غير مرتبط بموزّع بعينه (مثال: "12020001-ذمم لدى قسم التوزيع")
                    current_subaccount = narration
                    current_person = ""
            continue
        if debit == 0 and credit == 0:
            continue

        kind, key = extract_statement_key(narration)
        rows.append({
            "voucher_no": _clean(get(row, "رقم السند")),
            "narration": narration,
            "debit": debit,
            "credit": credit,
            "date": _clean(get(row, "التاريخ")),
            "subaccount": current_subaccount,
            "person": current_person,
            "kind": kind,
            "key": key,
            "statement_numbers": key if kind == "number" else set(),
        })
    return rows


# --- مطابقة تقريبية لأسماء الكشوفات (حين لا يوجد رقم) — نفس فكرة مطابقة
# أسماء الموزّعين بوحدة عمولات الموزّعين (طلب المستخدم صراحة 2026-09-29):
# اختلاف بسيط بحرف أو لاحقة موقع لا يمنع التجميع تحت نفس المجموعة. ---
def _norm_name(s):
    s = (s or "").replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    s = re.sub(r"[\.,]", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


class NameKeyResolver:
    """يوحّد أسماء كشوفات متقاربة الإملاء ضمن مجموعة واحدة، بنفس أسلوب
    NameResolver في وحدة عمولات الموزّعين: تطابق حرفي بعد التطبيع، ثم تطابق
    مجموعة كلمات فرعية، ثم أقرب تطابق (difflib) بحد أدنى 0.72."""

    def __init__(self, cutoff=0.72):
        self.cutoff = cutoff
        self.canonical = []  # قائمة الأسماء المعتمدة كما ظهرت أول مرة

    def resolve(self, raw):
        raw = (raw or "").strip()
        if not raw:
            return raw
        n = _norm_name(raw)
        for canon in self.canonical:
            if _norm_name(canon) == n:
                return canon
        n_tokens = set(n.split())
        if n_tokens:
            for canon in self.canonical:
                c_tokens = set(_norm_name(canon).split())
                if c_tokens and (n_tokens <= c_tokens or c_tokens <= n_tokens):
                    return canon
        normed = [_norm_name(c) for c in self.canonical]
        close = difflib.get_close_matches(n, normed, n=1, cutoff=self.cutoff)
        if close:
            for canon in self.canonical:
                if _norm_name(canon) == close[0]:
                    return canon
        self.canonical.append(raw)
        return raw


def _union_find_groups(rows):
    """يجمّع الحركات المرتبطة بنفس معرّف كشف ضمن مجموعة Union-Find واحدة.

    تصحيح 2026-09-29 (اكتُشف أثناء التحقق من الملف المرجعي الحقيقي، قبل أي
    عرض على المستخدم): معرّف "الرقم" مطابقته **عامة على مستوى الملف كله**
    عمداً — لأن تكرار نفس رقم الكشف تحت أكثر من حساب فرعي/موزّع هو بالضبط
    الآلية الحقيقية لاكتشاف المجموعات "المشتركة" (راجع توثيق أعلى الملف،
    مثال كشف 3310). أما معرّف "الاسم" (حين لا يوجد رقم) فمطابقته التقريبية
    (NameKeyResolver) يجب أن تكون **معزولة لكل شخص على حدة** فقط، تماماً
    كما في النموذج الأولي المُتحقَّق منه (full_proto.py) — لأن اسم صيدلية/جهة
    مرتبط بذمة موزّع بعينه، ومطابقة الأسماء التقريبية على مستوى الملف كله
    قد تدمج بالخطأ حركتين من موزّعين مختلفين لمجرد تشابه الاسم (تحقّق فعلي:
    الدمج العام على هذا الملف كان سيُنشئ 3 مجموعات "مشتركة" وهمية بلا أي
    أساس حقيقي، من بينها مثلاً اسم عام غير مرتبط بصيدلية بعينها مثل
    "ارصدة" — بخلاف حالات الرقم المشترك الحقيقية التي تحقّق منها المستخدم
    صراحة). لذلك: مطابقة الاسم تتم ضمن (الشخص، الاسم المُطبَّع) معاً، لا
    الاسم وحده."""
    parent = list(range(len(rows)))

    def find(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(a, b):
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[ra] = rb

    resolvers = {}  # شخص -> NameKeyResolver خاص به
    number_to_row = {}
    name_to_row = {}  # (شخص, الاسم المُطبَّع) -> فهرس الصف
    for i, r in enumerate(rows):
        if r["kind"] == "number":
            for num in r["key"]:
                if num in number_to_row:
                    union(i, number_to_row[num])
                else:
                    number_to_row[num] = i
        elif r["kind"] == "name":
            person = r["person"] or ""
            resolver = resolvers.setdefault(person, NameKeyResolver())
            resolved = resolver.resolve(r["key"])
            r["resolved_name"] = resolved
            key = (person, resolved)
            if key in name_to_row:
                union(i, name_to_row[key])
            else:
                name_to_row[key] = i

    groups = {}
    for i in range(len(rows)):
        root = find(i)
        groups.setdefault(root, []).append(i)
    return groups


def reconcile(rows, tolerance=Decimal("1")):
    linked = [r for r in rows if r["kind"] is not None]
    unlinked = [r for r in rows if r["kind"] is None]

    groups = _union_find_groups(linked)

    results = []
    for indices in groups.values():
        group_rows = [linked[i] for i in indices]
        numbers = sorted({n for r in group_rows if r["kind"] == "number" for n in r["key"]}, key=int)
        names = sorted({r.get("resolved_name", r["key"]) for r in group_rows if r["kind"] == "name"})
        persons = sorted({r["person"] for r in group_rows if r["person"]})
        total_debit = sum((r["debit"] for r in group_rows), Decimal("0"))
        total_credit = sum((r["credit"] for r in group_rows), Decimal("0"))
        diff = total_credit - total_debit
        matched = abs(diff) <= tolerance
        results.append({
            "statement_numbers": numbers,
            "statement_names": names,
            "rows": group_rows,
            "persons": persons,
            "is_shared": len(persons) > 1,
            "total_debit": total_debit,
            "total_credit": total_credit,
            "difference": diff,
            "matched": matched,
            "subaccounts": sorted({r["subaccount"] for r in group_rows if r["subaccount"]}),
        })

    def _sort_key(g):
        if g["statement_numbers"]:
            return (0, int(g["statement_numbers"][0]))
        return (1, g["statement_names"][0] if g["statement_names"] else "")

    results.sort(key=_sort_key)
    matched_groups = [g for g in results if g["matched"]]
    mismatched_groups = [g for g in results if not g["matched"]]

    # فروقات "فردية" (شخص واحد معني) مقابل "مشتركة" (أكثر من موزّع) — طلب
    # المستخدم الصريح 2026-09-29: تظهر المشتركة عند كل موزّع معني بوسم
    # "مشترك" دون تُحسَب في أي مجموع فردي.
    solo_mismatched = [g for g in mismatched_groups if not g["is_shared"]]
    shared_mismatched = [g for g in mismatched_groups if g["is_shared"]]

    summary = {
        "total_rows": len(rows),
        "linked_rows": len(linked),
        "unlinked_rows": len(unlinked),
        "groups_count": len(results),
        "matched_groups": len(matched_groups),
        "mismatched_groups": len(mismatched_groups),
        "solo_mismatched_groups": len(solo_mismatched),
        "shared_mismatched_groups": len(shared_mismatched),
        "mismatched_amount": sum((abs(g["difference"]) for g in mismatched_groups), Decimal("0")),
        "name_linked_groups": len([g for g in results if g["statement_names"] and not g["statement_numbers"]]),
    }
    return {
        "groups": results,
        "matched_groups": matched_groups,
        "mismatched_groups": mismatched_groups,
        "solo_mismatched_groups": solo_mismatched,
        "shared_mismatched_groups": shared_mismatched,
        "unlinked": unlinked,
        "summary": summary,
    }


# ------------------------------------------------- تفصيل حسب الشخص (بيد) ----
# طلب المستخدم صراحة 2026-09-29: "لازم يطلع عنا كل شخص والمبلغ اللي عليه فرق
# (+,-) ولازم يكون عنا الكشف الناقص أو الزائد ... يتم تعليمهم بحيث يكونو
# واضحين للقارئ". لكل موزّع: مجموع الفروقات "الفردية" (ناقص/زائد كل على حدة)
# + قائمة الكشوفات الناقصة والزائدة الخاصة به + قائمة الكشوفات المشتركة التي
# يظهر فيها (بلا احتساب في مجموعه الفردي) + حركاته بلا معرّف كشف على الإطلاق.
def compute_person_breakdown(result):
    persons = {}

    def _bucket(name):
        return persons.setdefault(name, {
            "name": name,
            "deficit_total": Decimal("0"),   # ناقص (مجموع الفروقات السالبة الفردية)
            "excess_total": Decimal("0"),    # زائد (مجموع الفروقات الموجبة الفردية)
            "deficit_groups": [],
            "excess_groups": [],
            "shared_groups": [],
            "unlinked_rows": [],
        })

    for g in result["solo_mismatched_groups"]:
        person = g["persons"][0] if g["persons"] else "بلا حساب فرعي معروف"
        b = _bucket(person)
        if g["difference"] < 0:
            b["deficit_total"] += g["difference"]
            b["deficit_groups"].append(g)
        else:
            b["excess_total"] += g["difference"]
            b["excess_groups"].append(g)

    for g in result["shared_mismatched_groups"]:
        for person in (g["persons"] or ["بلا حساب فرعي معروف"]):
            _bucket(person)["shared_groups"].append(g)

    for r in result["unlinked"]:
        person = r["person"] or "بلا حساب فرعي معروف"
        _bucket(person)["unlinked_rows"].append(r)

    for b in persons.values():
        b["net_total"] = b["deficit_total"] + b["excess_total"]

    return dict(sorted(persons.items()))
