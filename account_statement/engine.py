"""محرك فصل كشف حساب الصيدلية حسب العملة (الفكرة السابعة).

عمود "العملة الأصلية" في كشف الحساب يأتي بصيغة مثل:
    "16,848.00 [ل.س.ج. 1.00]"   → ليرة سورية جديدة
    "1,825,000.00 [ل.س.ق. 0.01]" → ليرة سورية قديمة
    "2.17 [$ 146.00]"            → دولار أمريكي

نستخرج رمز العملة من داخل الأقواس، ثم نقسّم حركات كشف الحساب حسب العملة،
مع مجموع مدين/دائن منفصل لكل عملة. يُبقي الكود على "الليرة القديمة" كفئة
مستقلة موضّحة (بدل دمجها قسراً مع الجديدة) حتى لا تُفقد أي بيانات صامتاً.
"""

import re
from collections import defaultdict
from decimal import Decimal

import openpyxl

CURRENCY_PATTERNS = [
    ("usd", re.compile(r"\$")),
    ("syp_old", re.compile(r"ل\.?\s*س\.?\s*ق")),
    ("syp_new", re.compile(r"ل\.?\s*س\.?\s*ج")),
]

CURRENCY_LABELS = {
    "usd": "دولار أمريكي",
    "syp_old": "ليرة سورية قديمة",
    "syp_new": "ليرة سورية جديدة",
    "unknown": "غير محدد",
}


def _clean(v):
    if v is None:
        return ""
    return str(v).strip() if isinstance(v, str) else v


def _to_decimal(v):
    if v is None or v == "":
        return Decimal("0")
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v)).quantize(Decimal("0.01"))
    s = str(v).replace(",", "").strip()
    if s in ("", "-"):
        return Decimal("0")
    try:
        return Decimal(s).quantize(Decimal("0.01"))
    except Exception:  # noqa: BLE001
        return Decimal("0")


def detect_currency(raw_value):
    text = str(raw_value) if raw_value is not None else ""
    for code, pattern in CURRENCY_PATTERNS:
        if pattern.search(text):
            return code
    return "unknown"


# صيغة عمود "العملة الأصلية" هي: "<القيمة بالعملة الأصلية> [<رمز العملة> <سعر الصرف>]"
# مثال: "43.40 [$ 146.00]" يعني 43.40 دولار بسعر صرف 146 (أي ما يعادل
# 43.40 × 146 = 6336.4 ل.س.ج، وهي بالضبط القيمة الظاهرة في عمود "مدين").
# أي أن عمودي "مدين"/"دائن" في الملف الخام هما دوماً القيمة المُحوَّلة
# لليرة السورية الجديدة (لأن الرصيد التراكمي في الحساب مُوحَّد بعملة واحدة)،
# وليسا القيمة الحقيقية بالعملة الأصلية للحركة. هذا هو سبب ظهور أرقام
# "غير صحيحة" سابقاً عند فصل الكشف حسب العملة: كانت الحركات بالدولار مثلاً
# تُعرض بقيمتها المحوَّلة لليرة (آلاف) بدل قيمتها الحقيقية بالدولار (عشرات).
_CURRENCY_CELL_RE = re.compile(r"^([\d,]+\.?\d*)\s*\[\s*(.+?)\s+([\d,]+\.?\d*)\s*\]")


def parse_currency_cell(raw_value):
    """يستخرج (القيمة الأصلية، سعر الصرف) من عمود "العملة الأصلية".
    يرجع (None, None) إن تعذّر التحليل (صيغة غير متوقعة)."""
    text = str(raw_value).strip() if raw_value is not None else ""
    m = _CURRENCY_CELL_RE.match(text)
    if not m:
        return None, None
    try:
        amount = Decimal(m.group(1).replace(",", ""))
        rate = Decimal(m.group(3).replace(",", ""))
    except Exception:  # noqa: BLE001
        return None, None
    return amount, rate


HEADER_SIGNATURE = {"رقم السند", "مدين", "دائن", "التاريخ"}
STOP_MARKERS = {"المجموع", "الرصيد النهائي"}
SKIP_MARKERS = {"الرصيد السابق"}


def parse_statement(file_obj):
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header_idx = None
    header = None
    for i, row in enumerate(all_rows):
        cells = {str(_clean(c)) for c in row if c is not None}
        if HEADER_SIGNATURE.issubset(cells):
            header_idx, header = i, [str(_clean(c)) for c in row]
            break
    if header_idx is None:
        raise ValueError("لم يتم العثور على صف العناوين المتوقع داخل كشف الحساب.")

    col = {name: header.index(name) for name in header if name}

    def get(row, name):
        idx = col.get(name)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    rows = []
    for row in all_rows[header_idx + 1:]:
        if row is None or all(c in (None, "") for c in row):
            continue
        first_cell = str(_clean(row[0])) if row[0] is not None else ""
        if first_cell in STOP_MARKERS:
            break
        if first_cell in SKIP_MARKERS:
            continue

        syp_debit = _to_decimal(get(row, "مدين"))
        syp_credit = _to_decimal(get(row, "دائن"))
        if syp_debit == 0 and syp_credit == 0:
            continue

        currency_raw = get(row, "العملة الأصلية")
        currency_code = detect_currency(currency_raw)
        fx_amount, fx_rate = parse_currency_cell(currency_raw)

        # مدين/دائن في الملف الخام هما دوماً القيمة المُحوَّلة لليرة السورية
        # الجديدة (لتوحيد الرصيد التراكمي بعملة واحدة)، وليسا القيمة الحقيقية
        # بالعملة الأصلية للحركة. إن أمكن استخراج القيمة الأصلية وسعر الصرف من
        # عمود "العملة الأصلية" نستخدمها هنا بدل القيمة المحوَّلة، على نفس
        # جهة الحركة (مدين/دائن) التي وردت بها أصلاً؛ وإلا نُبقي القيمة
        # المحوَّلة كما هي (عملة غير معروفة الصيغة).
        if fx_amount is not None:
            debit = fx_amount if syp_debit != 0 else Decimal("0")
            credit = fx_amount if syp_credit != 0 else Decimal("0")
        else:
            debit, credit = syp_debit, syp_credit

        rows.append({
            "voucher_no": _clean(get(row, "رقم السند")),
            "narration": _clean(get(row, "البيان")),
            "debit": debit,
            "credit": credit,
            "syp_equivalent_debit": syp_debit,
            "syp_equivalent_credit": syp_credit,
            "fx_rate": fx_rate,
            "date": _clean(get(row, "التاريخ")),
            "currency_raw": _clean(currency_raw),
            "currency_code": currency_code,
            "category": _clean(get(row, "الفئة")),
        })
    return rows


def split_by_currency(rows):
    buckets = defaultdict(list)
    for r in rows:
        buckets[r["currency_code"]].append(r)

    summary = {}
    for code, bucket_rows in buckets.items():
        summary[code] = {
            "label": CURRENCY_LABELS.get(code, code),
            "count": len(bucket_rows),
            "total_debit": sum((r["debit"] for r in bucket_rows), Decimal("0")),
            "total_credit": sum((r["credit"] for r in bucket_rows), Decimal("0")),
            # للمراجعة/التدقيق فقط: مجموع نفس الحركات مُحوَّلة لليرة السورية
            # الجديدة (كما وردت في عمودي مدين/دائن الخام)، للتأكد من أن مجموعها
            # يطابق إجمالي كشف الحساب الأصلي غير المُقسَّم.
            "total_debit_syp_equivalent": sum(
                (r.get("syp_equivalent_debit", r["debit"]) for r in bucket_rows), Decimal("0")
            ),
            "total_credit_syp_equivalent": sum(
                (r.get("syp_equivalent_credit", r["credit"]) for r in bucket_rows), Decimal("0")
            ),
        }
        summary[code]["balance"] = summary[code]["total_debit"] - summary[code]["total_credit"]

    return {"buckets": buckets, "summary": summary, "total_rows": len(rows)}
